from __future__ import annotations

from connection_status import cloud_device_status
from collector_update_api import register_update_routes, release_summary

import base64
import hashlib
import hmac
import io
import ipaddress
import json
import os
import re
import secrets
import threading
import time
import urllib.parse
import zipfile
from contextlib import nullcontext
from pathlib import Path
from typing import Any, Literal
from uuid import UUID, uuid4

from fastapi import (
    Cookie,
    Depends,
    FastAPI,
    Header,
    HTTPException,
    Query,
    Request,
    Response,
)
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, model_validator

from alarm_config import DEFAULT_THRESHOLDS, normalise_alarm_thresholds
from auth_service import AuthService
from email_alerts import enqueue_alarm
from email_api import register_email_routes
from diagnostics_api import register_diagnostic_routes
from report_i18n import (
    normalize_report_locale,
    report_alarm_details,
    report_catalog,
    report_metric_name,
    report_particle_unit,
    report_status,
)


def validate_particle_unit_metadata(
    particle_unit_code: int | None,
    particle_unit_label: str | None,
) -> None:
    if particle_unit_code is None and particle_unit_label is None:
        return  # Backward compatibility for records created before unit metadata existed.
    if particle_unit_code != 1 or particle_unit_label != "PCS/28.3L":
        raise ValueError("Current particle alarm policies require unit code 1 (PCS/28.3L)")


def authorize_ingest_target(
    db: Any,
    identity: dict[str, str],
    cleanroom_id: str,
    device_id: str,
    historical_at: float | None = None,
) -> tuple[str, str]:
    if historical_at is not None:
        row = db.execute(
            """SELECT cleanroom_name,device_name FROM device_assignment_periods
               WHERE customer_id=%s AND device_id=%s AND site_id=%s AND cleanroom_id=%s
                 AND started_at<=to_timestamp(%s) AND (ended_at IS NULL OR ended_at>=to_timestamp(%s))
               ORDER BY started_at DESC LIMIT 1""",
            (identity["customer_id"],device_id,identity["site_id"],cleanroom_id,historical_at,historical_at),
        ).fetchone()
    else:
        row = db.execute(
            """SELECT r.name,d.name FROM devices d JOIN cleanrooms r ON r.id=d.cleanroom_id
               WHERE d.id=%s AND d.customer_id=%s AND d.site_id=%s AND d.cleanroom_id=%s AND r.customer_id=%s AND d.enabled""",
            (device_id,identity["customer_id"],identity["site_id"],cleanroom_id,identity["customer_id"]),
        ).fetchone()
    if row is None:
        raise HTTPException(status_code=403, detail="Device is not assigned to this collector site")
    return str(row[0]), str(row[1])


class ReadingIn(BaseModel):
    record_uuid: UUID
    customer_id: str
    site_id: str
    cleanroom_id: str
    cleanroom_name: str
    device_id: str
    device_name: str
    measured_at: float = Field(gt=0)
    source: Literal["device"]
    particles: dict[str, int | float]
    environment: dict[str, int | float]
    alarm_status: str
    alarm_details: list[dict[str, Any]] = Field(default_factory=list)
    cleanliness_code: int | None = None
    cleanliness_label: str | None = None
    particle_unit_code: int | None = Field(default=None, ge=0, le=3)
    particle_unit_label: str | None = Field(default=None, max_length=32)
    protocol_profile: str | None = Field(default=None, max_length=120)

    @model_validator(mode="after")
    def validate_particle_unit(self) -> ReadingIn:
        validate_particle_unit_metadata(self.particle_unit_code, self.particle_unit_label)
        return self


class ReadingBatch(BaseModel):
    site_id: str
    readings: list[ReadingIn] = Field(min_length=1, max_length=1000)


class LatestReadingIn(BaseModel):
    customer_id: str
    site_id: str
    cleanroom_id: str
    cleanroom_name: str
    device_id: str
    device_name: str
    measured_at: float = Field(gt=0)
    source: Literal["device"]
    particles: dict[str, int | float]
    environment: dict[str, int | float]
    alarm_status: str
    alarm_details: list[dict[str, Any]] = Field(default_factory=list)
    cleanliness_code: int | None = None
    cleanliness_label: str | None = None
    particle_unit_code: int | None = Field(default=None, ge=0, le=3)
    particle_unit_label: str | None = Field(default=None, max_length=32)
    protocol_profile: str | None = Field(default=None, max_length=120)

    @model_validator(mode="after")
    def validate_particle_unit(self) -> LatestReadingIn:
        validate_particle_unit_metadata(self.particle_unit_code, self.particle_unit_label)
        return self


class LatestBatch(BaseModel):
    site_id: str
    readings: list[LatestReadingIn] = Field(min_length=1, max_length=1000)


class AlarmEventIn(BaseModel):
    event_uuid: UUID
    customer_id: str
    site_id: str
    cleanroom_id: str
    device_id: str
    source: Literal["device"]
    metric: str = Field(min_length=1, max_length=100)
    started_at: float = Field(gt=0, allow_inf_nan=False)
    ended_at: float | None = Field(default=None, gt=0, allow_inf_nan=False)
    trigger_value: float | None = Field(default=None, allow_inf_nan=False)
    peak_value: float | None = Field(default=None, allow_inf_nan=False)
    limit_description: str = Field(max_length=1000)

    @model_validator(mode="after")
    def validate_event_times(self) -> AlarmEventIn:
        if self.ended_at is not None and self.ended_at < self.started_at:
            raise ValueError("Alarm recovery time cannot precede its start")
        return self


class AlarmBatch(BaseModel):
    site_id: str
    events: list[AlarmEventIn] = Field(min_length=1, max_length=1000)


class LoginRequest(BaseModel):
    username: str = Field(min_length=1, max_length=200)
    password: str = Field(min_length=1, max_length=500)


class CustomerSettings(BaseModel):
    rooms: list[dict[str, Any]] = Field(min_length=1, max_length=200)


class CleanroomCreate(BaseModel):
    name: str = Field(min_length=1, max_length=100)


class DeviceCreate(BaseModel):
    cleanroom_id: str = Field(min_length=1, max_length=200)
    site_id: str | None = Field(default=None, max_length=200)
    name: str = Field(min_length=1, max_length=100)
    host: str = Field(min_length=1, max_length=64)
    tcp_port: int = Field(default=502, ge=1, le=65535)
    slave: int = Field(default=1, ge=1, le=247)


class CollectorCreate(BaseModel):
    name: str = Field(min_length=1, max_length=100)


class CollectorActivationClaim(BaseModel):
    activation_token: str = Field(min_length=24, max_length=500)
    machine_id: str = Field(pattern=r"^[0-9a-f]{32}$")
    hostname: str = Field(min_length=1, max_length=200)
    platform: str = Field(min_length=1, max_length=200)


class CollectorEnrollmentClaim(BaseModel):
    customer_id: str = Field(min_length=1, max_length=200)
    enrollment_token: str = Field(min_length=32, max_length=500)
    machine_id: str = Field(pattern=r"^[0-9a-f]{32}$")
    claim_secret: str = Field(min_length=32, max_length=500)
    hostname: str = Field(min_length=1, max_length=200)
    platform: str = Field(min_length=1, max_length=200)


class CollectorEnrollmentApproval(BaseModel):
    name: str = Field(min_length=1, max_length=100)


class CollectorPackageRequest(BaseModel):
    site_id: str = Field(min_length=1, max_length=200)
    activation_token: str = Field(min_length=24, max_length=500)


class DeviceConnectionStatus(BaseModel):
    device_id: str = Field(min_length=1, max_length=200)
    state: Literal["online", "offline", "unknown"]
    checked_at: float | None = Field(default=None, gt=0, allow_inf_nan=False)


class CollectorUpdateStatus(BaseModel):
    enabled: bool = False
    state: Literal["manual_upgrade_required", "starting", "restarting", "idle", "checking", "available", "downloading", "installing", "updated", "up_to_date", "failed", "rolled_back"] = "manual_upgrade_required"
    current_version: str | None = Field(default=None, max_length=80)
    available_version: str | None = Field(default=None, max_length=80)
    blocked_version: str | None = Field(default=None, max_length=80)
    updated_at: float | None = Field(default=None, ge=0, allow_inf_nan=False)
    last_check_at: float | None = Field(default=None, ge=0, allow_inf_nan=False)
    last_success_at: float | None = Field(default=None, ge=0, allow_inf_nan=False)
    error: str | None = Field(default=None, max_length=500)


class CollectorHeartbeat(BaseModel):
    site_id: str = Field(min_length=1, max_length=200)
    collector_time: float = Field(gt=0)
    version: str = Field(min_length=1, max_length=80)
    hostname: str = Field(min_length=1, max_length=200)
    platform: str = Field(min_length=1, max_length=200)
    monitor_running: bool
    device_total: int = Field(ge=0, le=10000)
    device_online: int = Field(ge=0, le=10000)
    update_status: CollectorUpdateStatus | None = None
    device_states: list[DeviceConnectionStatus] | None = Field(default=None, max_length=10000)
    last_reading_at: float | None = Field(default=None, gt=0)
    pending_uploads: int = Field(default=0, ge=0)
    quarantined_uploads: int = Field(default=0, ge=0)
    unassigned_uploads: int = Field(default=0, ge=0)
    oldest_pending_at: float | None = Field(default=None, gt=0)
    storage_state: str = Field(default="unknown", max_length=40)
    database_bytes: int | None = Field(default=None, ge=0)
    disk_free_bytes: int | None = Field(default=None, ge=0)
    applied_config_version: int | None = Field(default=None, ge=0)
    config_apply_status: Literal["awaiting", "applied", "failed"] = "awaiting"
    config_apply_error: str | None = Field(default=None, max_length=500)
    last_error: str | None = Field(default=None, max_length=500)

    @model_validator(mode="after")
    def validate_status_relationships(self) -> CollectorHeartbeat:
        if self.device_states is not None:
            ids = {item.device_id for item in self.device_states}
            if len(ids) != len(self.device_states) or len(ids) != self.device_total:
                raise ValueError("Device connection report must contain each device exactly once")
            if sum(item.state == "online" for item in self.device_states) != self.device_online:
                raise ValueError("Device connection counts do not match")
        if self.device_online > self.device_total:
            raise ValueError("Online device count exceeds total devices")
        if self.config_apply_status == "applied" and self.applied_config_version is None:
            raise ValueError("Applied configuration status requires an applied revision")
        return self


class DiscoveryRequest(BaseModel):
    cleanroom_id: str = Field(min_length=1, max_length=200)
    cidr: str | None = Field(default=None, max_length=64)
    tcp_port: int = Field(default=502, ge=1, le=65535)


class DiscoveryResultIn(BaseModel):
    host: str = Field(min_length=1, max_length=64)
    tcp_port: int = Field(ge=1, le=65535)
    slave: int = Field(default=1, ge=1, le=247)
    verified: bool = False
    modbus_responded: bool = False
    protocol_compatible: bool = False
    identity_verified: bool = False
    firmware_raw: int | None = Field(default=None, ge=0, le=65535)
    particle_unit_code: int | None = Field(default=None, ge=0, le=65535)
    particle_unit_label: str | None = Field(default=None, max_length=32)
    unit_supported: bool = False
    latency_ms: int | None = Field(default=None, ge=0, le=60000)

    @model_validator(mode="after")
    def validate_evidence(self) -> DiscoveryResultIn:
        expected_unit_label = {
            0: "PCS/L", 1: "PCS/28.3L", 2: "PCS/m3", 3: "PCS/2.83L",
        }.get(self.particle_unit_code)
        if self.particle_unit_label is not None and self.particle_unit_label != expected_unit_label:
            raise ValueError("Particle unit label does not match its register code")
        if self.unit_supported != (self.particle_unit_code == 1):
            raise ValueError("Unit support flag does not match particle unit code")
        if self.verified and not (
            self.modbus_responded and self.protocol_compatible and self.unit_supported
        ):
            raise ValueError("Verified discovery requires protocol and supported-unit evidence")
        return self


class DiscoveryCompletion(BaseModel):
    status: str = Field(pattern="^(completed|failed)$")
    scanned_cidr: str | None = Field(default=None, max_length=64)
    scanned_cidrs: list[str] = Field(default_factory=list, max_length=32)
    results: list[DiscoveryResultIn] = Field(default_factory=list, max_length=254)
    error: str | None = Field(default=None, max_length=500)


class DiscoveredDeviceAssignment(BaseModel):
    site_id: str = Field(min_length=1, max_length=200)
    cleanroom_id: str = Field(min_length=1, max_length=200)
    name: str = Field(min_length=1, max_length=100)
    host: str = Field(min_length=1, max_length=64)
    tcp_port: int = Field(default=502, ge=1, le=65535)
    slave: int = Field(default=1, ge=1, le=247)


def bounded_env_int(name: str, default: int, minimum: int, maximum: int) -> int:
    try:
        return max(minimum, min(int(os.environ.get(name, default)), maximum))
    except ValueError:
        return default


LOGIN_WINDOW_SECONDS = 15 * 60
LOGIN_MAX_FAILURES = 5
MAX_ACTIVE_DEVICES = bounded_env_int("DCP_MAX_ACTIVE_DEVICES", 100, 1, 1000)
MAX_CLEANROOMS = bounded_env_int("DCP_MAX_CLEANROOMS", 200, 1, 1000)
MAX_COLLECTORS = bounded_env_int("DCP_MAX_COLLECTORS", 100, 1, 1000)
MAX_DISCOVERY_ADDRESSES = 254
COLLECTOR_LEASE_SECONDS = 90
ACTIVATION_TTL_HOURS = bounded_env_int("DCP_ACTIVATION_TTL_HOURS", 24, 1, 168)
ENROLLMENT_TTL_HOURS = bounded_env_int("DCP_ENROLLMENT_TTL_HOURS", 24, 1, 168)
AUTO_DISCOVERY_INTERVAL_SECONDS = bounded_env_int(
    "DCP_AUTO_DISCOVERY_INTERVAL_SECONDS", 300, 60, 86400
)
DISCOVERY_EVIDENCE_TTL_SECONDS = bounded_env_int(
    "DCP_DISCOVERY_EVIDENCE_TTL_SECONDS", 900, 60, 86400
)
DISCOVERY_HISTORY_RETENTION_DAYS = bounded_env_int(
    "DCP_DISCOVERY_HISTORY_RETENTION_DAYS", 7, 1, 90
)
COLLECTOR_INSTANCE_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._:-]{7,127}$")
TOPOLOGY_MANAGER_ROLES = {"admin", "customer_admin", "customer"}
DATA_DESTRUCTION_ROLES = {"admin", "customer_admin"}
PRIVATE_DEVICE_NETWORKS = tuple(
    ipaddress.ip_network(value) for value in ("10.0.0.0/8", "172.16.0.0/12", "192.168.0.0/16")
)
_login_failures: dict[str, list[float]] = {}
_login_lock = threading.Lock()


def activation_secret_bytes() -> bytes:
    secret = os.environ.get("DCP_ACTIVATION_SECRET", "").strip()
    if not re.fullmatch(r"[0-9a-fA-F]{64}", secret):
        raise HTTPException(status_code=503, detail="Collector activation is not configured")
    return bytes.fromhex(secret)


def customer_enrollment_token(customer_id: str, generation: int) -> str:
    digest = hmac.new(
        activation_secret_bytes(),
        f"collector-enrollment:{customer_id}:{generation}".encode("utf-8"),
        hashlib.sha256,
    ).digest()
    return base64.urlsafe_b64encode(digest).rstrip(b"=").decode("ascii")


def enrolled_edge_token(customer_id: str, enrollment_id: str, machine_id: str) -> str:
    digest = hmac.new(
        activation_secret_bytes(),
        f"collector-edge:{customer_id}:{enrollment_id}:{machine_id}".encode("utf-8"),
        hashlib.sha256,
    ).digest()
    return base64.urlsafe_b64encode(digest).rstrip(b"=").decode("ascii")


def parse_scope_parameter(raw: str | None) -> list[str]:
    values: list[str] = []
    for item in str(raw or "").split(","):
        value = item.strip()
        if value and value not in values:
            values.append(value)
    if len(values) > MAX_ACTIVE_DEVICES or any(len(value) > 128 for value in values):
        raise HTTPException(status_code=400, detail="Too many or invalid scope identifiers")
    return values


def database_url() -> str:
    value = os.environ.get("DATABASE_URL", "").strip()
    if not value:
        raise RuntimeError("DATABASE_URL is required")
    return value


def collector_public_url(request: Request) -> str:
    configured = os.environ.get("DCP_PUBLIC_URL", "").strip().rstrip("/")
    value = configured or str(request.base_url).strip().rstrip("/")
    parsed = urllib.parse.urlparse(value)
    local_http = (
        parsed.scheme == "http"
        and parsed.hostname in {"127.0.0.1", "localhost", "::1"}
        and os.environ.get("DASHBOARD_SECURE_COOKIE", "1") == "0"
    )
    if not parsed.hostname or (parsed.scheme != "https" and not local_http):
        raise HTTPException(
            status_code=503,
            detail="DCP_PUBLIC_URL must be configured as the public HTTPS dashboard URL",
        )
    return value


def connect():
    try:
        import psycopg
    except ImportError as exc:
        raise RuntimeError("Install requirements-cloud.txt") from exc
    return psycopg.connect(database_url())


def initialize_schema() -> None:
    schema = Path(__file__).with_name("cloud_schema.sql").read_text(encoding="utf-8")
    with connect() as db:
        db.execute(schema)


def edge_identity(
    authorization: str | None = Header(default=None),
    x_dcp_collector_instance: str | None = Header(
        default=None, alias="X-DCP-Collector-Instance"
    ),
) -> dict[str, str]:
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Missing edge token")
    raw = authorization[7:].strip()
    token_hash = hashlib.sha256(raw.encode("utf-8")).hexdigest()
    instance_id = (x_dcp_collector_instance or f"legacy:{token_hash[:24]}").strip()
    if not COLLECTOR_INSTANCE_PATTERN.fullmatch(instance_id):
        raise HTTPException(status_code=400, detail="Invalid collector instance id")
    with connect() as db:
        row = db.execute(
            """
            SELECT customer_id,site_id FROM edge_tokens
            WHERE token_hash=%s AND enabled=true
            """,
            (token_hash,),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=401, detail="Invalid edge token")
        lease = db.execute(
            """
            UPDATE sites
            SET active_collector_instance=%s,active_collector_seen_at=now(),updated_at=now()
            WHERE id=%s AND customer_id=%s
              AND (
                active_collector_instance IS NULL
                OR active_collector_instance=%s
                OR active_collector_seen_at IS NULL
                OR active_collector_seen_at < now() - (%s * interval '1 second')
              )
            RETURNING id
            """,
            (instance_id, row[1], row[0], instance_id, COLLECTOR_LEASE_SECONDS),
        ).fetchone()
        if lease is None:
            raise HTTPException(
                status_code=409,
                detail="Another collector instance is already active for this node",
            )
        db.execute(
            "UPDATE edge_tokens SET last_used_at=now() WHERE token_hash=%s",
            (token_hash,),
        )
    return {
        "customer_id": str(row[0]), "site_id": str(row[1]),
        "instance_id": instance_id,
    }


def public_user(row: tuple[Any, ...]) -> dict[str, object]:
    return {
        "id": str(row[0]),
        "customer_id": row[1],
        "username": row[2],
        "display_name": row[3],
        "role": row[4],
    }


def customer_user(dcp_session: str | None = Cookie(default=None)) -> dict[str, object]:
    if not dcp_session:
        raise HTTPException(status_code=401, detail="Authentication required")
    token_hash = hashlib.sha256(dcp_session.encode("utf-8")).hexdigest()
    with connect() as db:
        row = db.execute(
            """
            SELECT users.id, users.customer_id, users.username, users.display_name, users.role
            FROM customer_sessions sessions
            JOIN customer_users users ON users.id = sessions.user_id
            WHERE sessions.token_hash = %s AND sessions.expires_at > now() AND users.enabled = true
            """,
            (token_hash,),
        ).fetchone()
    if row is None:
        raise HTTPException(status_code=401, detail="Authentication required")
    return public_user(row)


def topology_manager(user: dict[str, object] = Depends(customer_user)) -> dict[str, object]:
    """Authorize tenant topology changes; `customer` remains a legacy manager role."""
    if str(user.get("role", "")).casefold() not in TOPOLOGY_MANAGER_ROLES:
        raise HTTPException(status_code=403, detail="Administrator permission is required")
    return user


def is_primary_data_manager(user: dict[str, object], *, db: Any = None) -> bool:
    """Return whether this is the tenant's primary enabled administrator."""
    if str(user.get("role", "")).casefold() not in DATA_DESTRUCTION_ROLES:
        return False
    with (nullcontext(db) if db is not None else connect()) as connection:
        row = connection.execute(
            """
            SELECT id FROM customer_users
            WHERE customer_id=%s AND enabled=true AND lower(role)=ANY(%s)
            ORDER BY created_at,id LIMIT 1
            """,
            (str(user["customer_id"]), list(DATA_DESTRUCTION_ROLES)),
        ).fetchone()
    return bool(row and str(row[0]) == str(user.get("id")))


def primary_data_manager(
    user: dict[str, object] = Depends(customer_user),
) -> dict[str, object]:
    if not is_primary_data_manager(user):
        raise HTTPException(
            status_code=403,
            detail="Primary administrator permission is required for permanent deletion",
        )
    return user


def customer_configuration(
    customer_id: str, *, include_disabled: bool = False, db: Any = None,
) -> list[dict[str, object]]:
    with (nullcontext(db) if db is not None else connect()) as db:
        rooms = db.execute(
            "SELECT id, name FROM cleanrooms WHERE customer_id = %s ORDER BY sort_order, name",
            (customer_id,),
        ).fetchall()
        result = []
        for room_id, room_name in rooms:
            devices = []
            for row in db.execute(
                    f"""
                    SELECT device.id, device.name, device.enabled, device.host,
                        device.tcp_port, device.slave, device.site_id,
                        extract(epoch FROM max(latest.measured_at)) AS last_seen_at,
                        extract(epoch FROM device.disabled_at)
                    FROM devices AS device
                    LEFT JOIN latest_readings AS latest
                      ON latest.customer_id = device.customer_id
                     AND latest.device_id = device.id
                    WHERE device.customer_id = %s AND device.cleanroom_id = %s
                      {'' if include_disabled else 'AND device.enabled = true'}
                    GROUP BY device.id, device.name, device.enabled, device.host,
                        device.tcp_port, device.slave, device.site_id, device.sort_order, device.disabled_at
                    ORDER BY device.sort_order, device.name
                    """,
                    (customer_id, room_id),
                ).fetchall():
                last_seen_at = float(row[7]) if row[7] is not None else None
                devices.append(
                    {
                        "id": row[0], "name": row[1], "enabled": row[2], "host": row[3],
                        "tcpPort": row[4], "slave": row[5], "site_id": row[6],
                        "last_seen_at": last_seen_at,
                        "disabled_at": float(row[8]) if row[8] is not None else None,
                        "online": bool(row[2] and last_seen_at and time.time() - last_seen_at <= 45),
                    }
                )
            for device in devices:
                maintenance = db.execute(
                    """SELECT reason,extract(epoch FROM ends_at) FROM device_maintenance_periods
                       WHERE customer_id=%s AND device_id=%s AND started_at<=now()
                         AND LEAST(ends_at,COALESCE(ended_at,ends_at))>now()
                       ORDER BY started_at DESC LIMIT 1""", (customer_id,device["id"]),
                ).fetchone()
                device["maintenance"] = {"reason":maintenance[0],"until":float(maintenance[1])} if maintenance else None
            if include_disabled:
                known = {d["id"] for d in devices}
                for old in db.execute(
                    """SELECT DISTINCT ON (p.device_id) p.device_id,p.device_name,d.enabled,extract(epoch FROM d.disabled_at)
                       FROM device_assignment_periods p JOIN devices d ON d.id=p.device_id
                       WHERE p.customer_id=%s AND p.cleanroom_id=%s AND d.cleanroom_id<>p.cleanroom_id
                       ORDER BY p.device_id,p.started_at DESC""", (customer_id,room_id),
                ).fetchall():
                    if old[0] not in known:
                        devices.append({"id":old[0],"name":old[1],"enabled":old[2],"historical_assignment":True,
                                        "disabled_at":float(old[3]) if old[3] else None,"online":False})
            threshold = db.execute(
                """
                SELECT profile_name,
                    particle_0_3_max, particle_0_3_enabled,
                    particle_0_5_max, particle_0_5_enabled,
                    particle_1_0_max, particle_1_0_enabled,
                    particle_2_5_max, particle_2_5_enabled,
                    particle_5_0_max, particle_5_0_enabled,
                    particle_10_0_max, particle_10_0_enabled,
                    temperature_min, temperature_max, humidity_min, humidity_max,
                    alarm_delay_seconds
                FROM thresholds WHERE customer_id = %s AND cleanroom_id = %s
                """,
                (customer_id, room_id),
            ).fetchone()
            result.append(
                {
                    "id": room_id,
                    "name": room_name,
                    "devices": devices,
                    "thresholds": {
                        "profile_name": threshold[0],
                        "particle_0_3_max": threshold[1],
                        "particle_0_3_enabled": threshold[2],
                        "particle_0_5_max": threshold[3],
                        "particle_0_5_enabled": threshold[4],
                        "particle_1_0_max": threshold[5],
                        "particle_1_0_enabled": threshold[6],
                        "particle_2_5_max": threshold[7],
                        "particle_2_5_enabled": threshold[8],
                        "particle_5_0_max": threshold[9],
                        "particle_5_0_enabled": threshold[10],
                        "particle_10_0_max": threshold[11],
                        "particle_10_0_enabled": threshold[12],
                        "temperature_min": threshold[13],
                        "temperature_max": threshold[14],
                        "humidity_min": threshold[15],
                        "humidity_max": threshold[16],
                        "alarm_delay_seconds": threshold[17],
                    } if threshold else None,
                }
            )
    return result


def configuration_for_site(
    rooms: list[dict[str, object]], site_id: str,
) -> list[dict[str, object]]:
    """Return every customer room but only the devices owned by one edge site."""
    return [
        {
            **room,
            "devices": [
                device for device in room.get("devices", [])
                if isinstance(device, dict) and str(device.get("site_id", "")) == site_id
            ],
        }
        for room in rooms
    ]


def validated_device_connection(device: dict[str, Any]) -> tuple[str, str, int, int]:
    name = str(device.get("name", "")).strip()
    host = str(device.get("host", "")).strip()
    if not name or len(name) > 100:
        raise HTTPException(status_code=400, detail="Device name is required and must be 100 characters or fewer")
    try:
        address = ipaddress.ip_address(host)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"Invalid device IP address: {host}") from exc
    if address.version != 4 or not any(address in network for network in PRIVATE_DEVICE_NETWORKS):
        raise HTTPException(status_code=400, detail="Device IP must be a private IPv4 address")
    try:
        tcp_port = int(device.get("tcpPort", 502))
        slave = int(device.get("slave", 1))
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Port and Slave ID must be whole numbers") from exc
    if not 1 <= tcp_port <= 65535:
        raise HTTPException(status_code=400, detail="Device port must be between 1 and 65535")
    if not 1 <= slave <= 247:
        raise HTTPException(status_code=400, detail="Slave ID must be between 1 and 247")
    return name, str(address), tcp_port, slave


def insert_default_cloud_thresholds(db: Any, customer_id: str, room_id: str) -> None:
    values = normalise_alarm_thresholds(DEFAULT_THRESHOLDS)
    db.execute(
        """
        INSERT INTO thresholds(
            cleanroom_id,customer_id,profile_name,
            particle_0_3_max,particle_0_3_enabled,
            particle_0_5_max,particle_0_5_enabled,
            particle_1_0_max,particle_1_0_enabled,
            particle_2_5_max,particle_2_5_enabled,
            particle_5_0_max,particle_5_0_enabled,
            particle_10_0_max,particle_10_0_enabled,particle_5_max,
            temperature_min,temperature_max,humidity_min,humidity_max,
            alarm_delay_seconds
        ) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """,
        (
            room_id, customer_id, values["profile_name"],
            values["particle_0_3_max"], values["particle_0_3_enabled"],
            values["particle_0_5_max"], values["particle_0_5_enabled"],
            values["particle_1_0_max"], values["particle_1_0_enabled"],
            values["particle_2_5_max"], values["particle_2_5_enabled"],
            values["particle_5_0_max"], values["particle_5_0_enabled"],
            values["particle_10_0_max"], values["particle_10_0_enabled"],
            values["particle_0_5_max"],
            values["temperature_min"], values["temperature_max"],
            values["humidity_min"], values["humidity_max"],
            values["alarm_delay_seconds"],
        ),
    )


def record_configuration_audit(
    db: Any,
    customer_id: str,
    actor_kind: str,
    action: str,
    target_type: str,
    target_id: str,
    details: dict[str, object],
    actor_user_id: str | None = None,
) -> None:
    db.execute(
        """
        INSERT INTO configuration_audit_events(
            id,customer_id,actor_user_id,actor_kind,action,target_type,target_id,details
        ) VALUES(%s,%s,%s,%s,%s,%s,%s,%s::jsonb)
        """,
        (
            uuid4(), customer_id, actor_user_id, actor_kind, action,
            target_type, target_id, json.dumps(details, ensure_ascii=False),
        ),
    )


def create_cloud_collector(
    customer_id: str,
    name: object,
    *,
    actor_user_id: str,
) -> dict[str, object]:
    """Create one isolated collector node; package issuance is a separate retryable step."""
    collector_name = str(name or "").strip()
    if not collector_name or len(collector_name) > 100:
        raise HTTPException(
            status_code=400,
            detail="Collector name is required and must be 100 characters or fewer",
        )
    site_id = f"collector-{secrets.token_hex(12)}"
    with connect() as db:
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR UPDATE", (customer_id,))
        if db.execute(
            "SELECT 1 FROM sites WHERE customer_id=%s AND lower(name)=lower(%s)",
            (customer_id, collector_name),
        ).fetchone():
            raise HTTPException(status_code=409, detail="Collector name is already in use")
        count = int(db.execute(
            "SELECT COUNT(*) FROM sites WHERE customer_id=%s", (customer_id,)
        ).fetchone()[0])
        if count >= MAX_COLLECTORS:
            raise HTTPException(
                status_code=400,
                detail=f"A customer can have at most {MAX_COLLECTORS} collectors",
            )
        db.execute(
            "INSERT INTO sites(id,customer_id,name) VALUES(%s,%s,%s)",
            (site_id, customer_id, collector_name),
        )
        record_configuration_audit(
            db, customer_id, "customer_user", "collector.created", "collector", site_id,
            {
                "name": collector_name,
                "package_ready": True,
            }, actor_user_id,
        )
    return {
        "collector": {
            "id": site_id,
            "name": collector_name,
            "connected": False,
            "connection_state": "awaiting_activation",
            "config_version": 1,
            "applied_config_version": None,
            "config_state": "awaiting",
        },
        "package": {"site_id": site_id, "ready": True},
    }


def claim_collector_activation(
    activation_token: str,
    machine_id: str,
    hostname: str,
    platform_name: str,
) -> dict[str, str]:
    """Claim once, but return the same credential to the same installation on retry."""
    token_hash = hashlib.sha256(activation_token.encode("utf-8")).hexdigest()
    edge_token = base64.urlsafe_b64encode(
        hmac.new(
            activation_secret_bytes(),
            f"{token_hash}:{machine_id}".encode("ascii"),
            hashlib.sha256,
        ).digest()
    ).rstrip(b"=").decode("ascii")
    edge_hash = hashlib.sha256(edge_token.encode("utf-8")).hexdigest()
    with connect() as db:
        row = db.execute(
            """
            SELECT customer_id,site_id,used_at,claimed_machine_id
            FROM collector_activation_tokens
            WHERE token_hash=%s AND expires_at > now()
            FOR UPDATE
            """,
            (token_hash,),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=401, detail="Activation code is invalid or expired")
        customer_id, site_id = str(row[0]), str(row[1])
        first_claim = row[2] is None
        if not first_claim and str(row[3] or "") != machine_id:
            raise HTTPException(status_code=401, detail="Activation code has already been used")
        if first_claim:
            db.execute(
                """
                UPDATE collector_activation_tokens
                SET used_at=now(),claimed_machine_id=%s
                WHERE token_hash=%s
                """,
                (machine_id, token_hash),
            )
            db.execute(
                """
                INSERT INTO edge_tokens(token_hash,customer_id,site_id,label)
                VALUES(%s,%s,%s,%s)
                """,
                (edge_hash, customer_id, site_id, f"Activated by {hostname}"[:200]),
            )
            record_configuration_audit(
                db, customer_id, "edge", "collector.activated", "collector", site_id,
                {"hostname": hostname, "platform": platform_name, "machine_id": machine_id},
            )
        elif db.execute(
            """
            SELECT 1 FROM edge_tokens
            WHERE token_hash=%s AND customer_id=%s AND site_id=%s AND enabled=true
            """,
            (edge_hash, customer_id, site_id),
        ).fetchone() is None:
            raise HTTPException(status_code=401, detail="Activated collector credential is disabled")
    return {"site_id": site_id, "token": edge_token}


def register_collector_enrollment(
    customer_id: str,
    enrollment_token: str,
    machine_id: str,
    claim_secret: str,
    hostname: str,
    platform_name: str,
) -> dict[str, object]:
    """Register one machine without granting it an upload credential until approval."""
    claim_hash = hashlib.sha256(claim_secret.encode("utf-8")).hexdigest()
    with connect() as db:
        customer = db.execute(
            "SELECT enrollment_generation FROM customers WHERE id=%s FOR SHARE",
            (customer_id,),
        ).fetchone()
        if customer is None:
            raise HTTPException(status_code=401, detail="Collector enrollment is invalid")
        expected = customer_enrollment_token(customer_id, int(customer[0]))
        if not hmac.compare_digest(expected, enrollment_token):
            raise HTTPException(status_code=401, detail="Collector enrollment is invalid")
        # Serialize registrations per tenant so two first requests cannot create
        # duplicate machine rows or collide on the legacy internal short code.
        db.execute(
            "SELECT pg_advisory_xact_lock(hashtextextended(%s,0))",
            (f"collector-enrollment:{customer_id}",),
        )
        row = db.execute(
            """
            SELECT id,claim_secret_hash,pairing_code,status,site_id,expires_at,
                   enrollment_generation
            FROM collector_enrollments
            WHERE customer_id=%s AND machine_id=%s
            FOR UPDATE
            """,
            (customer_id, machine_id),
        ).fetchone()
        if (
            row is not None
            and str(row[3]) != "revoked"
            and not hmac.compare_digest(str(row[1]), claim_hash)
        ):
            raise HTTPException(status_code=401, detail="Collector enrollment is invalid")
        if row is not None and str(row[3]) == "approved":
            enrollment_id, site_id = str(row[0]), str(row[4] or "")
            edge_token = enrolled_edge_token(customer_id, enrollment_id, machine_id)
            edge_hash = hashlib.sha256(edge_token.encode("utf-8")).hexdigest()
            if not site_id or db.execute(
                """
                SELECT 1 FROM edge_tokens
                WHERE token_hash=%s AND customer_id=%s AND site_id=%s AND enabled=true
                """,
                (edge_hash, customer_id, site_id),
            ).fetchone() is None:
                raise HTTPException(status_code=403, detail="Collector enrollment is disabled")
            db.execute(
                "UPDATE collector_enrollments SET last_seen_at=now() WHERE id=%s",
                (row[0],),
            )
            return {"status": "approved", "site_id": site_id, "token": edge_token}
        if row is not None and str(row[3]) == "revoked":
            if int(row[6]) >= int(customer[0]):
                raise HTTPException(status_code=403, detail="Collector enrollment is disabled")
            for _attempt in range(8):
                pairing_code = secrets.token_hex(4).upper()
                if db.execute(
                    """
                    SELECT 1 FROM collector_enrollments
                    WHERE customer_id=%s AND pairing_code=%s AND status='pending'
                    """,
                    (customer_id, pairing_code),
                ).fetchone() is None:
                    break
            else:
                raise HTTPException(status_code=503, detail="Could not allocate a pairing code")
            db.execute(
                """
                UPDATE collector_enrollments
                SET claim_secret_hash=%s,pairing_code=%s,status='pending',site_id=NULL,
                    hostname=%s,platform=%s,enrollment_generation=%s,
                    last_seen_at=now(),expires_at=now()+(%s * interval '1 hour'),
                    approved_at=NULL,approved_by=NULL
                WHERE id=%s
                """,
                (
                    claim_hash, pairing_code, hostname, platform_name, int(customer[0]),
                    ENROLLMENT_TTL_HOURS, row[0],
                ),
            )
            return {
                "status": "pending",
                "expires_in_seconds": ENROLLMENT_TTL_HOURS * 3600,
            }
        if row is None:
            pending_count = int(db.execute(
                """
                SELECT COUNT(*) FROM collector_enrollments
                WHERE customer_id=%s AND status='pending' AND expires_at > now()
                """,
                (customer_id,),
            ).fetchone()[0])
            if pending_count >= MAX_COLLECTORS * 2:
                raise HTTPException(status_code=429, detail="Too many pending collectors")
            enrollment_id = uuid4()
            for _attempt in range(8):
                pairing_code = secrets.token_hex(4).upper()
                if db.execute(
                    """
                    SELECT 1 FROM collector_enrollments
                    WHERE customer_id=%s AND pairing_code=%s AND status='pending'
                    """,
                    (customer_id, pairing_code),
                ).fetchone() is None:
                    break
            else:
                raise HTTPException(status_code=503, detail="Could not allocate a pairing code")
            db.execute(
                """
                INSERT INTO collector_enrollments(
                    id,customer_id,machine_id,enrollment_generation,
                    claim_secret_hash,pairing_code,
                    hostname,platform,expires_at
                ) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,now()+(%s * interval '1 hour'))
                """,
                (
                    enrollment_id, customer_id, machine_id, int(customer[0]),
                    claim_hash, pairing_code,
                    hostname, platform_name, ENROLLMENT_TTL_HOURS,
                ),
            )
            record_configuration_audit(
                db, customer_id, "edge", "collector.enrollment_requested",
                "collector_enrollment", str(enrollment_id),
                {"hostname": hostname, "platform": platform_name, "machine_id": machine_id},
            )
        else:
            enrollment_id, pairing_code = row[0], str(row[2])
            db.execute(
                """
                UPDATE collector_enrollments
                SET hostname=%s,platform=%s,last_seen_at=now(),
                    expires_at=now()+(%s * interval '1 hour')
                WHERE id=%s
                """,
                (hostname, platform_name, ENROLLMENT_TTL_HOURS, enrollment_id),
            )
        return {
            "status": "pending",
            "expires_in_seconds": ENROLLMENT_TTL_HOURS * 3600,
        }


def approve_collector_enrollment(
    customer_id: str,
    enrollment_id: str,
    name: object,
    *,
    actor_user_id: str,
) -> dict[str, object]:
    collector_name = str(name or "").strip()
    if not collector_name or len(collector_name) > 100:
        raise HTTPException(status_code=400, detail="Collector name is required")
    try:
        parsed_enrollment_id = UUID(enrollment_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail="Pending collector not found") from exc
    with connect() as db:
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR UPDATE", (customer_id,))
        row = db.execute(
            """
            SELECT machine_id,hostname,platform,status
            FROM collector_enrollments
            WHERE id=%s AND customer_id=%s AND expires_at > now()
            FOR UPDATE
            """,
            (parsed_enrollment_id, customer_id),
        ).fetchone()
        if row is None or str(row[3]) != "pending":
            raise HTTPException(status_code=409, detail="Pending collector is missing or expired")
        if db.execute(
            "SELECT 1 FROM sites WHERE customer_id=%s AND lower(name)=lower(%s)",
            (customer_id, collector_name),
        ).fetchone():
            raise HTTPException(status_code=409, detail="Collector name is already in use")
        count = int(db.execute(
            "SELECT COUNT(*) FROM sites WHERE customer_id=%s", (customer_id,)
        ).fetchone()[0])
        if count >= MAX_COLLECTORS:
            raise HTTPException(
                status_code=400,
                detail=f"A customer can have at most {MAX_COLLECTORS} collectors",
            )
        site_id = f"collector-{secrets.token_hex(12)}"
        edge_token = enrolled_edge_token(customer_id, str(parsed_enrollment_id), str(row[0]))
        edge_hash = hashlib.sha256(edge_token.encode("utf-8")).hexdigest()
        db.execute(
            "INSERT INTO sites(id,customer_id,name) VALUES(%s,%s,%s)",
            (site_id, customer_id, collector_name),
        )
        db.execute(
            """
            INSERT INTO edge_tokens(token_hash,customer_id,site_id,label)
            VALUES(%s,%s,%s,%s)
            """,
            (edge_hash, customer_id, site_id, f"Enrolled by {row[1]}"[:200]),
        )
        db.execute(
            """
            UPDATE collector_enrollments
            SET status='approved',site_id=%s,approved_at=now(),approved_by=%s
            WHERE id=%s
            """,
            (site_id, actor_user_id, parsed_enrollment_id),
        )
        record_configuration_audit(
            db, customer_id, "customer_user", "collector.enrollment_approved",
            "collector", site_id,
            {
                "name": collector_name, "hostname": str(row[1]),
                "platform": str(row[2]),
            },
            actor_user_id,
        )
    return {
        "id": site_id,
        "name": collector_name,
        "connected": False,
        "connection_state": "awaiting_activation",
    }


def revoke_collector_enrollment(
    customer_id: str, enrollment_id: str, *, actor_user_id: str
) -> None:
    try:
        parsed_enrollment_id = UUID(enrollment_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail="Pending collector not found") from exc
    with connect() as db:
        row = db.execute(
            """
            UPDATE collector_enrollments SET status='revoked'
            WHERE id=%s AND customer_id=%s AND status='pending'
            RETURNING machine_id
            """,
            (parsed_enrollment_id, customer_id),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Pending collector not found")
        record_configuration_audit(
            db, customer_id, "customer_user", "collector.enrollment_rejected",
            "collector_enrollment", str(parsed_enrollment_id),
            {"machine_id": str(row[0])}, actor_user_id,
        )


def reissue_collector_activation(
    customer_id: str,
    site_id: str,
    *,
    actor_user_id: str,
) -> str:
    """Issue a fresh claim only while a collector has never been activated."""
    activation_token = secrets.token_urlsafe(32)
    activation_hash = hashlib.sha256(activation_token.encode("utf-8")).hexdigest()
    with connect() as db:
        site = db.execute(
            "SELECT name FROM sites WHERE id=%s AND customer_id=%s FOR UPDATE",
            (site_id, customer_id),
        ).fetchone()
        if site is None:
            raise HTTPException(status_code=404, detail="Collector site not found")
        if db.execute(
            "SELECT 1 FROM edge_tokens WHERE customer_id=%s AND site_id=%s AND enabled=true LIMIT 1",
            (customer_id, site_id),
        ).fetchone():
            raise HTTPException(
                status_code=409,
                detail="Collector is already activated; a second package cannot be issued",
            )
        db.execute(
            """
            DELETE FROM collector_activation_tokens
            WHERE customer_id=%s AND site_id=%s
              AND created_at < now() - interval '30 days'
            """,
            (customer_id, site_id),
        )
        # A retry must be safe after a failed browser download or an expired
        # package. Only the newest unconsumed package remains usable.
        db.execute(
            "DELETE FROM collector_activation_tokens WHERE customer_id=%s AND site_id=%s AND used_at IS NULL",
            (customer_id, site_id),
        )
        db.execute(
            """
            INSERT INTO collector_activation_tokens(
                token_hash,customer_id,site_id,created_by,expires_at
            ) VALUES(%s,%s,%s,%s,now()+(%s * interval '1 hour'))
            """,
            (
                activation_hash, customer_id, site_id, actor_user_id,
                ACTIVATION_TTL_HOURS,
            ),
        )
        record_configuration_audit(
            db, customer_id, "customer_user", "collector.package_issued",
            "collector", site_id,
            {
                "activation_ttl_hours": ACTIVATION_TTL_HOURS,
                "previous_unconsumed_packages_revoked": True,
            },
            actor_user_id,
        )
    return activation_token


def published_collector_executable() -> bytes:
    """Load a release EXE only when its independently published digest matches."""
    exe_path = Path(
        os.environ.get(
            "DCP_COLLECTOR_EXE_PATH",
            str(Path(__file__).with_name("release") / "HawkHive-DPC8001-Collector.exe"),
        )
    ).resolve()
    if not exe_path.is_file():
        raise HTTPException(status_code=503, detail="Windows collector package is not published")
    exe = exe_path.read_bytes()
    if len(exe) > 100 * 1024 * 1024 or not exe.startswith(b"MZ"):
        raise HTTPException(status_code=503, detail="Published Windows collector is invalid")
    checksum_path = exe_path.with_suffix(exe_path.suffix + ".sha256")
    try:
        expected_hash = checksum_path.read_text(encoding="ascii").split()[0].lower()
    except (OSError, IndexError, UnicodeError):
        raise HTTPException(status_code=503, detail="Windows collector checksum is not published")
    if not hmac.compare_digest(expected_hash, hashlib.sha256(exe).hexdigest()):
        raise HTTPException(status_code=503, detail="Published Windows collector checksum is invalid")
    return exe


def build_collector_package(
    customer_id: str,
    site_id: str,
    activation_token: str,
    cloud_url: str,
) -> bytes:
    """Build a customer-bound ZIP without exposing a durable edge credential."""
    activation_hash = hashlib.sha256(activation_token.encode("utf-8")).hexdigest()
    with connect() as db:
        row = db.execute(
            """
            SELECT 1 FROM collector_activation_tokens
            WHERE token_hash=%s AND customer_id=%s AND site_id=%s
              AND used_at IS NULL AND expires_at > now()
            """,
            (activation_hash, customer_id, site_id),
        ).fetchone()
    if row is None:
        raise HTTPException(status_code=409, detail="Activation package is invalid or expired")
    exe = published_collector_executable()
    activation = json.dumps(
        {
            "version": 1,
            "cloud_url": cloud_url.rstrip("/"),
            "activation_token": activation_token,
        },
        separators=(",", ":"),
    ).encode("utf-8")
    instructions = (
        "HawkHive 自动采集器\r\n\r\n"
        "1. 解压整个文件夹。\r\n"
        "2. 双击 HawkHive-DPC8001-Collector.exe，并在管理员权限提示中选择“是”。\r\n"
        "3. 程序会自动安装为开机后台任务；电脑不需要安装 Python。\r\n"
        "4. 首次启动会自动连接云端并删除一次性激活文件。\r\n"
        "5. 等待设备出现在云端“待分配设备”，再选择车间并修改名称。\r\n"
        "6. 不要把本文件夹复制给其他客户或其他采集电脑。\r\n"
    ).encode("utf-8-sig")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        archive.writestr("HawkHive-Collector/HawkHive-DPC8001-Collector.exe", exe)
        archive.writestr("HawkHive-Collector/hawkhive-activation.json", activation)
        archive.writestr("HawkHive-Collector/安装说明.txt", instructions)
    return output.getvalue()


def build_reusable_collector_package(customer_id: str, cloud_url: str) -> bytes:
    """Build one customer-scoped package that may enroll any number of PCs."""
    with connect() as db:
        row = db.execute(
            "SELECT enrollment_generation FROM customers WHERE id=%s",
            (customer_id,),
        ).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Customer not found")
    exe = published_collector_executable()
    enrollment = json.dumps(
        {
            "version": 1,
            "cloud_url": cloud_url.rstrip("/"),
            "customer_id": customer_id,
            "enrollment_token": customer_enrollment_token(customer_id, int(row[0])),
        },
        separators=(",", ":"),
    ).encode("utf-8")
    instructions = (
        "HawkHive 通用自动采集器（本客户专用）\r\n\r\n"
        "1. 本压缩包可复制到本客户的多台 Windows 电脑重复安装。\r\n"
        "2. 每台电脑解压整个文件夹后，双击 HawkHive-DPC8001-Collector.exe。\r\n"
        "3. 在管理员权限提示中选择“是”；电脑不需要安装 Python。\r\n"
        "4. 安装完成后，电脑会自动出现在云端的待加入列表。\r\n"
        "5. 云端管理员决定加入并填写采集器名称后，电脑会自动取得独立凭证并开始工作。\r\n"
        "6. 同一台电脑可负责多个车间；设备通过后再在云端分配车间和名称。\r\n"
        "7. 本包只能用于当前客户，请勿发送给其他客户。\r\n"
    ).encode("utf-8-sig")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        archive.writestr("HawkHive-Collector/HawkHive-DPC8001-Collector.exe", exe)
        archive.writestr("HawkHive-Collector/hawkhive-enrollment.json", enrollment)
        archive.writestr("HawkHive-Collector/安装说明.txt", instructions)
    return output.getvalue()


def rotate_customer_enrollment(customer_id: str, *, actor_user_id: str) -> int:
    """Invalidate copied enrollment packages without touching approved collectors."""
    with connect() as db:
        row = db.execute(
            """
            UPDATE customers SET enrollment_generation=enrollment_generation+1
            WHERE id=%s RETURNING enrollment_generation
            """,
            (customer_id,),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Customer not found")
        revoked = db.execute(
            """
            UPDATE collector_enrollments SET status='revoked'
            WHERE customer_id=%s AND status='pending'
            RETURNING id
            """,
            (customer_id,),
        ).fetchall()
        record_configuration_audit(
            db, customer_id, "customer_user", "collector.enrollment_rotated",
            "customer", customer_id,
            {"generation": int(row[0]), "pending_revoked": len(revoked)},
            actor_user_id,
        )
    return int(row[0])


def create_cloud_cleanroom(
    customer_id: str,
    name: object,
    *,
    actor_kind: str,
    actor_user_id: str | None = None,
) -> str:
    room_name = str(name or "").strip()
    if not room_name or len(room_name) > 100:
        raise HTTPException(
            status_code=400,
            detail="Cleanroom name is required and must be 100 characters or fewer",
        )
    room_id = f"room-{secrets.token_hex(16)}"
    with connect() as db:
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR UPDATE", (customer_id,))
        if db.execute(
            "SELECT 1 FROM cleanrooms WHERE customer_id=%s AND lower(name)=lower(%s)",
            (customer_id, room_name),
        ).fetchone():
            raise HTTPException(status_code=409, detail="Cleanroom name is already in use")
        count = int(db.execute(
            "SELECT COUNT(*) FROM cleanrooms WHERE customer_id=%s", (customer_id,)
        ).fetchone()[0])
        if count >= MAX_CLEANROOMS:
            raise HTTPException(
                status_code=400,
                detail=f"A customer can have at most {MAX_CLEANROOMS} cleanrooms",
            )
        sort_order = int(db.execute(
            "SELECT COALESCE(MAX(sort_order),0)+1 FROM cleanrooms WHERE customer_id=%s",
            (customer_id,),
        ).fetchone()[0])
        db.execute(
            "INSERT INTO cleanrooms(id,customer_id,name,sort_order) VALUES(%s,%s,%s,%s)",
            (room_id, customer_id, room_name, sort_order),
        )
        insert_default_cloud_thresholds(db, customer_id, room_id)
        db.execute(
            "UPDATE sites SET config_version=config_version+1 WHERE customer_id=%s",
            (customer_id,),
        )
        record_configuration_audit(
            db, customer_id, actor_kind, "cleanroom.created", "cleanroom", room_id,
            {"name": room_name}, actor_user_id,
        )
    return room_id


def create_cloud_device(
    customer_id: str,
    payload: dict[str, Any],
    *,
    actor_kind: str,
    actor_user_id: str | None = None,
    forced_site_id: str | None = None,
) -> str:
    cleanroom_id = str(payload.get("cleanroom_id", "")).strip()
    device_name, host, tcp_port, slave = validated_device_connection(payload)
    requested_site_id = str(forced_site_id or payload.get("site_id") or "").strip()
    with connect() as db:
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR UPDATE", (customer_id,))
        room = db.execute(
            "SELECT 1 FROM cleanrooms WHERE id=%s AND customer_id=%s",
            (cleanroom_id, customer_id),
        ).fetchone()
        if room is None:
            raise HTTPException(status_code=404, detail="Cleanroom not found")
        if requested_site_id:
            site = db.execute(
                "SELECT id FROM sites WHERE id=%s AND customer_id=%s",
                (requested_site_id, customer_id),
            ).fetchone()
        else:
            sites = db.execute(
                "SELECT id FROM sites WHERE customer_id=%s ORDER BY created_at,id LIMIT 2",
                (customer_id,),
            ).fetchall()
            if not sites:
                raise HTTPException(status_code=400, detail="No site collector is configured")
            if len(sites) > 1:
                raise HTTPException(status_code=400, detail="Select a site collector for this device")
            site = sites[0]
        if site is None:
            raise HTTPException(status_code=404, detail="Site collector not found")
        site_id = str(site[0])
        if db.execute(
            """
            SELECT 1 FROM devices
            WHERE customer_id=%s AND cleanroom_id=%s AND enabled=true AND lower(name)=lower(%s)
            """,
            (customer_id, cleanroom_id, device_name),
        ).fetchone():
            raise HTTPException(status_code=409, detail="Device name is already in use in this cleanroom")
        if db.execute(
            """
            SELECT 1 FROM devices
            WHERE customer_id=%s AND site_id=%s AND enabled=true
              AND host=%s AND tcp_port=%s AND slave=%s
            """,
            (customer_id, site_id, host, tcp_port, slave),
        ).fetchone():
            raise HTTPException(
                status_code=409,
                detail=f"Device connection {host}:{tcp_port} / slave {slave} is already configured",
            )
        count = int(db.execute(
            "SELECT COUNT(*) FROM devices WHERE customer_id=%s AND enabled=true",
            (customer_id,),
        ).fetchone()[0])
        if count >= MAX_ACTIVE_DEVICES:
            raise HTTPException(
                status_code=400,
                detail=f"A customer can have at most {MAX_ACTIVE_DEVICES} active devices",
            )
        sort_order = int(db.execute(
            """
            SELECT COALESCE(MAX(sort_order),0)+1 FROM devices
            WHERE customer_id=%s AND cleanroom_id=%s AND enabled=true
            """,
            (customer_id, cleanroom_id),
        ).fetchone()[0])
        device_id = f"device-{secrets.token_hex(16)}"
        db.execute(
            """
            INSERT INTO devices(
                id,customer_id,cleanroom_id,site_id,name,host,tcp_port,slave,sort_order
            ) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)
            """,
            (
                device_id, customer_id, cleanroom_id, site_id, device_name,
                host, tcp_port, slave, sort_order,
            ),
        )
        db.execute(
            "UPDATE sites SET config_version=config_version+1 WHERE id=%s AND customer_id=%s",
            (site_id, customer_id),
        )
        record_configuration_audit(
            db, customer_id, actor_kind, "device.created", "device", device_id,
            {
                "name": device_name, "cleanroom_id": cleanroom_id, "site_id": site_id,
                "host": host, "tcp_port": tcp_port, "slave": slave,
                "state": "pending_collector_confirmation",
            }, actor_user_id,
        )
    return device_id


def validated_discovery_cidr(value: str | None) -> str | None:
    if value is None or not value.strip():
        return None
    try:
        network = ipaddress.ip_network(value.strip(), strict=False)
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail="Scan network must be a valid IPv4 CIDR, for example 192.168.1.0/24",
        ) from exc
    if network.version != 4 or not any(network.subnet_of(allowed) for allowed in PRIVATE_DEVICE_NETWORKS):
        raise HTTPException(status_code=400, detail="Scan network must be inside a private IPv4 range")
    if network.num_addresses > MAX_DISCOVERY_ADDRESSES + 2:
        raise HTTPException(status_code=400, detail="Scan network may contain at most 254 usable addresses")
    return str(network)


DISCOVERY_JOB_COLUMNS = """
id,status,requested_cidr,scanned_cidr,tcp_port,requested_at,started_at,completed_at,
result_count,error,cleanroom_id,site_id,scanned_cidrs,source
"""


def discovery_job_dict(
    row: tuple[Any, ...], results: list[dict[str, object]] | None = None
) -> dict[str, object]:
    return {
        "id": str(row[0]), "status": row[1], "requested_cidr": row[2],
        "scanned_cidr": row[3], "tcp_port": row[4],
        "requested_at": row[5].timestamp(),
        "started_at": row[6].timestamp() if row[6] else None,
        "completed_at": row[7].timestamp() if row[7] else None,
        "result_count": row[8], "error": row[9], "cleanroom_id": row[10],
        "site_id": row[11], "scanned_cidrs": row[12] or [],
        "source": row[13], "results": results or [],
    }


def reading_dict(row: tuple[Any, ...]) -> dict[str, object]:
    measured_at = row[5].timestamp()
    return {
        "record_uuid": str(row[0]) if row[0] else None, "cleanroom_id": row[1], "cleanroom": row[2],
        "device_id": row[3], "device": row[4], "timestamp": measured_at,
        "source": row[6], "particles": row[7], "environment": row[8],
        "alarm_status": row[9], "alarm_details": row[10],
        "cleanliness_code": row[11], "cleanliness_label": row[12],
        "particle_unit_code": row[13], "particle_unit_label": row[14],
        "protocol_profile": row[15],
        "online": time.time() - measured_at <= 45,
        # Cloud freshness is not evidence of a broken device-to-collector link.
        "connection_state": "online" if time.time() - measured_at <= 45 else "sync_stale",
    }


def annotate_maintenance(records: list[dict], customer_id: str, *, current: bool = False) -> list[dict]:
    if not records:
        return records
    with connect() as db:
        periods = db.execute(
            """SELECT device_id,reason,extract(epoch FROM started_at),
                      extract(epoch FROM LEAST(ends_at,COALESCE(ended_at,ends_at)))
               FROM device_maintenance_periods WHERE customer_id=%s AND device_id=ANY(%s)
               ORDER BY started_at DESC""", (customer_id,list({str(r["device_id"]) for r in records})),
        ).fetchall()
    by_device: dict[str,list] = {}
    for period in periods:
        by_device.setdefault(period[0],[]).append(period)
    for record in records:
        stamp = time.time() if current else record.get("timestamp",record.get("started_at",0))
        period = next((p for p in by_device.get(record["device_id"],[]) if float(p[2])<=stamp<float(p[3])),None)
        record["maintenance"] = {"reason":period[1],"until":float(period[3])} if period else None
    return records


def apply_display_names(readings: list[dict[str, object]], customer_id: str) -> list[dict[str, object]]:
    config = customer_configuration(customer_id)
    room_names = {room["id"]: room["name"] for room in config}
    device_names = {
        device["id"]: device["name"] for room in config for device in room["devices"]
    }
    for reading in readings:
        reading["cleanroom"] = room_names.get(reading["cleanroom_id"], reading["cleanroom"])
        reading["device"] = device_names.get(reading["device_id"], reading["device"])
    return readings


_docs_enabled = os.environ.get("CLOUD_ENABLE_DOCS") == "1"
app = FastAPI(
    title="HawkHive Cloud Monitoring Platform", version="1.0.0",
    docs_url="/docs" if _docs_enabled else None,
    redoc_url="/redoc" if _docs_enabled else None,
    openapi_url="/openapi.json" if _docs_enabled else None,
)


@app.on_event("startup")
def startup() -> None:
    initialize_schema()


@app.get("/api/v1/health")
def health() -> dict[str, object]:
    with connect() as db:
        db.execute("SELECT 1").fetchone()
    return {"ok": True}


@app.post("/api/login")
def customer_login(payload: LoginRequest, request: Request, response: Response) -> dict[str, object]:
    key = f"{request.client.host if request.client else 'unknown'}:{payload.username.casefold()}"
    now = time.time()
    with _login_lock:
        failures = [stamp for stamp in _login_failures.get(key, []) if now - stamp < LOGIN_WINDOW_SECONDS]
        _login_failures[key] = failures
        if len(failures) >= LOGIN_MAX_FAILURES:
            raise HTTPException(status_code=429, detail="Too many login attempts; try again later")
    with connect() as db:
        row = db.execute(
            """
            SELECT id, customer_id, username, display_name, role, password_salt, password_hash
            FROM customer_users WHERE lower(username) = lower(%s) AND enabled = true
            """,
            (payload.username.strip(),),
        ).fetchone()
        if row is None or not AuthService.verify_password(payload.password, row[5], row[6]):
            with _login_lock:
                _login_failures.setdefault(key, []).append(now)
            raise HTTPException(status_code=401, detail="Invalid username or password")
        token = secrets.token_urlsafe(32)
        token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
        db.execute("DELETE FROM customer_sessions WHERE expires_at <= now()")
        db.execute(
            """
            INSERT INTO customer_sessions(token_hash, user_id, expires_at)
            VALUES (%s, %s, now() + interval '12 hours')
            """,
            (token_hash, row[0]),
        )
    with _login_lock:
        _login_failures.pop(key, None)
    response.set_cookie(
        "dcp_session", token, max_age=43200, httponly=True, samesite="lax",
        secure=os.environ.get("DASHBOARD_SECURE_COOKIE", "1") != "0", path="/",
    )
    return {"ok": True, "data": public_user(row[:5])}


@app.post("/api/logout")
def customer_logout(response: Response, dcp_session: str | None = Cookie(default=None)) -> dict[str, object]:
    if dcp_session:
        token_hash = hashlib.sha256(dcp_session.encode("utf-8")).hexdigest()
        with connect() as db:
            db.execute("DELETE FROM customer_sessions WHERE token_hash = %s", (token_hash,))
    response.delete_cookie("dcp_session", path="/")
    return {"ok": True}


@app.get("/api/session")
def customer_session(user: dict[str, object] = Depends(customer_user)) -> dict[str, object]:
    return {"ok": True, "data": {**user, "email_alerts_available": True}}


@app.post("/api/v1/edge/activate")
def activate_collector(payload: CollectorActivationClaim) -> dict[str, object]:
    return {
        "ok": True,
        "data": claim_collector_activation(
            payload.activation_token.strip(), payload.machine_id,
            payload.hostname.strip(), payload.platform.strip()
        ),
    }


@app.post("/api/v1/edge/enroll")
def enroll_collector(payload: CollectorEnrollmentClaim) -> dict[str, object]:
    return {
        "ok": True,
        "data": register_collector_enrollment(
            payload.customer_id.strip(), payload.enrollment_token.strip(),
            payload.machine_id, payload.claim_secret.strip(),
            payload.hostname.strip(), payload.platform.strip(),
        ),
    }


@app.get("/api/config")
def get_customer_config(
    user: dict[str, object] = Depends(customer_user), include_disabled: bool = False,
) -> dict[str, object]:
    return {"ok": True, "data": customer_configuration(str(user["customer_id"]), include_disabled=include_disabled)}


@app.get("/api/admin/sites")
def get_admin_sites(
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    with connect() as db:
        rows = db.execute(
            """
            SELECT
              site.id,site.name,site.config_version,site.last_heartbeat_at,
              site.applied_config_version,site.config_apply_status,
              site.config_apply_error,site.collector_status,token.last_used_at
            FROM sites AS site
            LEFT JOIN (
              SELECT customer_id,site_id,max(last_used_at) AS last_used_at
              FROM edge_tokens WHERE enabled=true
              GROUP BY customer_id,site_id
            ) AS token
              ON token.customer_id=site.customer_id AND token.site_id=site.id
            WHERE site.customer_id=%s
            ORDER BY site.created_at,site.id
            """,
            (customer_id,),
        ).fetchall()
        can_destroy_data = is_primary_data_manager(user, db=db)
    now = time.time()
    sites: list[dict[str, object]] = []
    for row in rows:
        heartbeat_at = row[3].timestamp() if row[3] else None
        last_contact_at = row[8].timestamp() if row[8] else None
        connected = bool(heartbeat_at and 0 <= now - heartbeat_at <= COLLECTOR_LEASE_SECONDS)
        legacy_contact = bool(
            not heartbeat_at and last_contact_at
            and now - last_contact_at <= COLLECTOR_LEASE_SECONDS
        )
        status = row[7] if isinstance(row[7], dict) else {}
        desired_revision = int(row[2])
        applied_revision = int(row[4]) if row[4] is not None else None
        apply_status = str(row[5] or "awaiting")
        if apply_status == "failed":
            config_state = "failed"
        elif applied_revision is not None and applied_revision >= desired_revision:
            config_state = "applied"
        elif heartbeat_at:
            config_state = "pending"
        else:
            config_state = "awaiting"
        sites.append({
            "id": str(row[0]), "name": str(row[1]),
            "config_version": desired_revision,
            "applied_config_version": applied_revision,
            "config_apply_status": apply_status,
            "config_apply_error": row[6],
            "config_state": config_state,
            "last_heartbeat_at": heartbeat_at,
            "last_contact_at": last_contact_at,
            "connected": connected,
            "connection_state": (
                "online" if connected else "upgrade_required" if legacy_contact
                else "offline" if heartbeat_at else "awaiting_activation"
            ),
            "legacy_contact": legacy_contact,
            "is_current": False,
            "version": status.get("version"),
            "update_status": status.get("update_status") or {"enabled": False, "state": "manual_upgrade_required"},
            "hostname": status.get("hostname"),
            "platform": status.get("platform"),
            "monitor_running": status.get("monitor_running"),
            **cloud_device_status(status, connected),
            "last_reading_at": status.get("last_reading_at"),
            "pending_uploads": int(status.get("pending_uploads") or 0),
            "quarantined_uploads": int(status.get("quarantined_uploads") or 0),
            "unassigned_uploads": int(status.get("unassigned_uploads") or 0),
            "oldest_pending_at": status.get("oldest_pending_at"),
            "storage_state": status.get("storage_state", "unknown"),
            "disk_free_bytes": status.get("disk_free_bytes"),
            "clock_offset_seconds": status.get("clock_offset_seconds"),
            "last_error": status.get("last_error"),
        })
    return {
        "ok": True,
        "data": {
            "sites": sites,
            "max_cleanrooms": MAX_CLEANROOMS,
            "max_active_devices": MAX_ACTIVE_DEVICES,
            "max_collectors": MAX_COLLECTORS,
            "manual_registration": True,
            "can_delete_devices": True,
            "can_delete_workshops": can_destroy_data,
            "can_cleanup_data": can_destroy_data,
            "device_lifecycle": True,
            "connection_verification": "collector",
            "management_scope": "customer",
            "can_create_collectors": True,
            "diagnostic_logs": True,
            "collector_update": release_summary(),
            "can_download_collector": True,
            "scope_notice": "可在这里管理本客户的全部采集器节点。",
        },
    }


@app.get("/api/admin/pending-collectors")
def admin_pending_collectors(
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    with connect() as db:
        rows = db.execute(
            """
            SELECT id,hostname,platform,
                   extract(epoch FROM created_at),extract(epoch FROM last_seen_at),
                   extract(epoch FROM expires_at)
            FROM collector_enrollments
            WHERE customer_id=%s AND status='pending' AND expires_at > now()
            ORDER BY last_seen_at DESC,id
            """,
            (str(user["customer_id"]),),
        ).fetchall()
    return {
        "ok": True,
        "data": [
            {
                "id": str(row[0]),
                "hostname": str(row[1]), "platform": str(row[2]),
                "created_at": float(row[3]), "last_seen_at": float(row[4]),
                "expires_at": float(row[5]),
            }
            for row in rows
        ],
    }


@app.post("/api/admin/pending-collectors/{enrollment_id}/approve", status_code=201)
def admin_approve_pending_collector(
    enrollment_id: str,
    payload: CollectorEnrollmentApproval,
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    return {
        "ok": True,
        "data": approve_collector_enrollment(
            str(user["customer_id"]), enrollment_id, payload.name,
            actor_user_id=str(user["id"]),
        ),
    }


@app.delete("/api/admin/pending-collectors/{enrollment_id}")
def admin_reject_pending_collector(
    enrollment_id: str,
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    revoke_collector_enrollment(
        str(user["customer_id"]), enrollment_id, actor_user_id=str(user["id"]),
    )
    return {"ok": True}


@app.post("/api/admin/collector-installer")
def admin_reusable_collector_installer(
    request: Request,
    user: dict[str, object] = Depends(topology_manager),
) -> StreamingResponse:
    package = build_reusable_collector_package(
        str(user["customer_id"]), collector_public_url(request),
    )
    return StreamingResponse(
        io.BytesIO(package),
        media_type="application/zip",
        headers={
            "Content-Disposition": 'attachment; filename="HawkHive-Collector-Windows.zip"',
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.post("/api/admin/collector-installer/rotate")
def admin_rotate_collector_installer(
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    generation = rotate_customer_enrollment(
        str(user["customer_id"]), actor_user_id=str(user["id"]),
    )
    return {"ok": True, "data": {"generation": generation}}


@app.post("/api/admin/collectors", status_code=201)
def admin_create_collector(
    payload: CollectorCreate,
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    return {
        "ok": True,
        "data": create_cloud_collector(
            str(user["customer_id"]), payload.name, actor_user_id=str(user["id"]),
        ),
    }


@app.post("/api/admin/collector-package")
def admin_collector_package(
    payload: CollectorPackageRequest,
    request: Request,
    user: dict[str, object] = Depends(topology_manager),
) -> StreamingResponse:
    package = build_collector_package(
        str(user["customer_id"]), payload.site_id,
        payload.activation_token.strip(), collector_public_url(request),
    )
    return StreamingResponse(
        io.BytesIO(package),
        media_type="application/zip",
        headers={
            "Content-Disposition": 'attachment; filename="HawkHive-Collector-Windows.zip"',
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.post("/api/admin/collectors/{site_id}/package")
def admin_reissue_collector_package(
    site_id: str,
    request: Request,
    user: dict[str, object] = Depends(topology_manager),
) -> StreamingResponse:
    customer_id = str(user["customer_id"])
    # Validate the release artifact before consuming/replacing an activation
    # package. A broken cloud image must not invalidate the administrator's
    # last downloadable package.
    public_url = collector_public_url(request)
    published_collector_executable()
    activation_token = reissue_collector_activation(
        customer_id, site_id, actor_user_id=str(user["id"]),
    )
    package = build_collector_package(
        customer_id, site_id, activation_token, public_url,
    )
    return StreamingResponse(
        io.BytesIO(package),
        media_type="application/zip",
        headers={
            "Content-Disposition": 'attachment; filename="HawkHive-Collector-Windows.zip"',
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.post("/api/admin/cleanrooms", status_code=201)
def admin_create_cleanroom(
    payload: CleanroomCreate,
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    create_cloud_cleanroom(
        customer_id, payload.name, actor_kind="customer_user",
        actor_user_id=str(user["id"]),
    )
    return {"ok": True, "data": customer_configuration(customer_id)}


@app.post("/api/admin/devices", status_code=201)
def admin_create_device(
    payload: DeviceCreate,
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    create_cloud_device(
        customer_id,
        {
            "cleanroom_id": payload.cleanroom_id, "site_id": payload.site_id,
            "name": payload.name, "host": payload.host,
            "tcp_port": payload.tcp_port, "slave": payload.slave,
        },
        actor_kind="customer_user", actor_user_id=str(user["id"]),
    )
    return {"ok": True, "data": customer_configuration(customer_id)}


@app.delete("/api/admin/devices/{device_id}")
def admin_delete_device(
    device_id: str,
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    """Retire a registration while retaining readings and alarm history."""
    customer_id = str(user["customer_id"])
    with connect() as db:
        # Serialize with registration/configuration changes for this customer.
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR UPDATE", (customer_id,))
        device = db.execute(
            """
            SELECT name,cleanroom_id,site_id,host,tcp_port,slave,enabled
            FROM devices WHERE id=%s AND customer_id=%s FOR UPDATE
            """,
            (device_id, customer_id),
        ).fetchone()
        if device is None:
            raise HTTPException(status_code=404, detail="Device not found")
        # A retried request must not create another audit event or revision.
        if device[6]:
            db.execute(
                """
                UPDATE devices SET enabled=false,disabled_at=clock_timestamp(),updated_at=now()
                WHERE id=%s AND customer_id=%s
                """,
                (device_id, customer_id),
            )
            db.execute(
                "UPDATE sites SET config_version=config_version+1 WHERE id=%s AND customer_id=%s",
                (device[2], customer_id),
            )
            record_configuration_audit(
                db, customer_id, "customer_user", "device.deleted", "device", device_id,
                {
                    "name": device[0], "cleanroom_id": device[1], "site_id": device[2],
                    "host": device[3], "tcp_port": device[4], "slave": device[5],
                    "history_retained": True,
                }, str(user["id"]),
            )
        updated = customer_configuration(customer_id, db=db)
    return {"ok": True, "data": updated}


@app.get("/api/admin/discovered-devices")
def admin_discovered_devices(
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    with connect() as db:
        rows = db.execute(
            """
            WITH latest AS (
                SELECT DISTINCT ON (job.site_id,result.host,result.tcp_port,result.slave)
                    job.site_id,result.host,result.tcp_port,result.slave,
                    result.firmware_raw,result.particle_unit_code,
                    result.particle_unit_label,result.latency_ms,result.discovered_at
                FROM device_discovery_results AS result
                JOIN device_discovery_jobs AS job ON job.id=result.job_id
                WHERE job.customer_id=%s AND job.status='completed'
                  AND result.verified=true AND result.protocol_compatible=true
                  AND result.unit_supported=true
                  AND result.discovered_at > now()-(%s * interval '1 second')
                ORDER BY job.site_id,result.host,result.tcp_port,result.slave,
                    result.discovered_at DESC
            )
            SELECT latest.site_id,site.name,latest.host,latest.tcp_port,latest.slave,
                latest.firmware_raw,latest.particle_unit_code,latest.particle_unit_label,
                latest.latency_ms,latest.discovered_at
            FROM latest
            JOIN sites AS site ON site.id=latest.site_id AND site.customer_id=%s
            WHERE NOT EXISTS (
                SELECT 1 FROM devices AS device
                WHERE device.customer_id=%s AND device.site_id=latest.site_id
                  AND device.host=latest.host AND device.tcp_port=latest.tcp_port
                  AND device.slave=latest.slave AND device.enabled=true
            )
            ORDER BY latest.discovered_at DESC,latest.site_id,latest.host
            """,
            (
                customer_id, DISCOVERY_EVIDENCE_TTL_SECONDS,
                customer_id, customer_id,
            ),
        ).fetchall()
    return {
        "ok": True,
        "data": [
            {
                "site_id": str(row[0]), "site_name": str(row[1]),
                "host": str(row[2]), "tcp_port": int(row[3]), "slave": int(row[4]),
                "firmware_raw": row[5], "particle_unit_code": row[6],
                "particle_unit_label": row[7], "latency_ms": row[8],
                "discovered_at": row[9].timestamp(),
            }
            for row in rows
        ],
    }


@app.post("/api/admin/discovered-devices/assign", status_code=201)
def admin_assign_discovered_device(
    payload: DiscoveredDeviceAssignment,
    user: dict[str, object] = Depends(topology_manager),
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    try:
        host = str(ipaddress.ip_address(payload.host.strip()))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Discovered device address is invalid") from exc
    with connect() as db:
        evidence = db.execute(
            """
            SELECT 1
            FROM device_discovery_results AS result
            JOIN device_discovery_jobs AS job ON job.id=result.job_id
            WHERE job.customer_id=%s AND job.site_id=%s AND job.status='completed'
              AND result.host=%s AND result.tcp_port=%s AND result.slave=%s
              AND result.verified=true AND result.protocol_compatible=true
              AND result.unit_supported=true
              AND result.discovered_at > now()-(%s * interval '1 second')
            ORDER BY result.discovered_at DESC LIMIT 1
            """,
            (
                customer_id, payload.site_id, host, payload.tcp_port, payload.slave,
                DISCOVERY_EVIDENCE_TTL_SECONDS,
            ),
        ).fetchone()
    if evidence is None:
        raise HTTPException(status_code=409, detail="Verified discovery evidence is no longer available")
    create_cloud_device(
        customer_id,
        {
            "cleanroom_id": payload.cleanroom_id,
            "site_id": payload.site_id,
            "name": payload.name,
            "host": host,
            "tcp_port": payload.tcp_port,
            "slave": payload.slave,
        },
        actor_kind="customer_user", actor_user_id=str(user["id"]),
    )
    return {"ok": True, "data": customer_configuration(customer_id)}


def apply_customer_config(
    payload: CustomerSettings,
    user: dict[str, object],
    *,
    allow_connection_changes: bool,
    forced_site_id: str | None = None,
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    complete_configuration = customer_configuration(customer_id)
    visible_configuration = (
        configuration_for_site(complete_configuration, forced_site_id)
        if forced_site_id else complete_configuration
    )
    known = {room["id"]: room for room in visible_configuration}
    if {str(room.get("id", "")) for room in payload.rooms} != set(known):
        raise HTTPException(status_code=400, detail="Cleanroom list cannot be changed by a customer")
    if not allow_connection_changes:
        safe_rooms: list[dict[str, Any]] = []
        for room in payload.rooms:
            room_id = str(room.get("id", ""))
            known_devices = {
                str(device["id"]): device for device in known[room_id]["devices"]
            }
            requested_devices = room.get("devices")
            if not isinstance(requested_devices, list):
                raise HTTPException(status_code=400, detail="Devices must be a list")
            requested_by_id = {
                str(device.get("id", "")): device
                for device in requested_devices
                if isinstance(device, dict)
            }
            if set(requested_by_id) != set(known_devices):
                raise HTTPException(
                    status_code=400,
                    detail="Device list cannot be changed by a customer",
                )
            safe_rooms.append({
                **room,
                "devices": [
                    {
                        **known_device,
                        "name": str(
                            requested_by_id[device_id].get(
                                "name", known_device["name"]
                            )
                        ).strip(),
                    }
                    for device_id, known_device in known_devices.items()
                ],
            })
        payload = CustomerSettings(rooms=safe_rooms)
    device_lists = [room.get("devices", []) for room in payload.rooms]
    if any(not isinstance(devices, list) for devices in device_lists):
        raise HTTPException(status_code=400, detail="Devices must be a list")
    if sum(len(devices) for devices in device_lists) > MAX_ACTIVE_DEVICES:
        raise HTTPException(status_code=400, detail=f"A customer can have at most {MAX_ACTIVE_DEVICES} active devices")
    room_names = [str(room.get("name", "")).strip().casefold() for room in payload.rooms]
    if any(not name for name in room_names) or len(room_names) != len(set(room_names)):
        raise HTTPException(status_code=400, detail="Cleanroom names must be non-empty and unique")
    seen_endpoints: set[tuple[str, int, int]] = set()
    seen_device_ids: set[str] = set()
    for room in payload.rooms:
        devices = room.get("devices", [])
        if not isinstance(devices, list):
            continue
        for device in devices:
            if not isinstance(device, dict):
                raise HTTPException(status_code=400, detail="Invalid device settings")
            requested_id = str(device.get("id", "")).strip()
            if requested_id:
                if requested_id in seen_device_ids:
                    raise HTTPException(status_code=400, detail="A device can appear only once in the configuration")
                seen_device_ids.add(requested_id)
            _device_name, host, tcp_port, slave = validated_device_connection(device)
            endpoint = (host, tcp_port, slave)
            if endpoint in seen_endpoints:
                raise HTTPException(
                    status_code=400,
                    detail=f"Device connection {host}:{tcp_port} / slave {slave} is already configured",
                )
            seen_endpoints.add(endpoint)

    with connect() as db:
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR UPDATE", (customer_id,))
        if forced_site_id:
            site = db.execute(
                "SELECT id FROM sites WHERE id=%s AND customer_id=%s",
                (forced_site_id, customer_id),
            ).fetchone()
        else:
            site = db.execute(
                "SELECT id FROM sites WHERE customer_id = %s ORDER BY created_at, id LIMIT 1",
                (customer_id,),
            ).fetchone()
        if site is None:
            raise HTTPException(status_code=400, detail="No site is configured for this customer")
        site_id = str(site[0])
        for room in payload.rooms:
            room_id = str(room["id"])
            name = str(room.get("name", "")).strip()
            db.execute(
                "UPDATE cleanrooms SET name = %s WHERE id = %s AND customer_id = %s",
                (name, room_id, customer_id),
            )
            known_devices = {str(item["id"]): item for item in known[room_id]["devices"]}
            devices = room.get("devices", [])
            if not isinstance(devices, list):
                raise HTTPException(status_code=400, detail="Devices must be a list")
            names: set[str] = set()
            retained: set[str] = set()
            for sort_order, device in enumerate(devices, 1):
                if not isinstance(device, dict):
                    raise HTTPException(status_code=400, detail="Invalid device settings")
                device_name, host, tcp_port, slave = validated_device_connection(device)
                folded = device_name.casefold()
                if folded in names:
                    raise HTTPException(status_code=400, detail="Device names must be unique within a cleanroom")
                names.add(folded)
                requested_id = str(device.get("id", "")).strip()
                if requested_id:
                    if requested_id not in known_devices:
                        raise HTTPException(status_code=400, detail="Unknown device")
                    retained.add(requested_id)
                    if forced_site_id:
                        db.execute(
                            """
                            UPDATE devices SET name=%s,host=%s,tcp_port=%s,slave=%s,
                                sort_order=%s,updated_at=now()
                            WHERE id=%s AND customer_id=%s AND cleanroom_id=%s
                              AND site_id=%s AND enabled=true
                            """,
                            (
                                device_name,host,tcp_port,slave,sort_order,requested_id,
                                customer_id,room_id,site_id,
                            ),
                        )
                    else:
                        db.execute(
                            """
                            UPDATE devices SET name=%s,host=%s,tcp_port=%s,slave=%s,
                                sort_order=%s,updated_at=now()
                            WHERE id=%s AND customer_id=%s AND cleanroom_id=%s AND enabled=true
                            """,
                            (
                                device_name,host,tcp_port,slave,sort_order,requested_id,
                                customer_id,room_id,
                            ),
                        )
                    previous = known_devices[requested_id]
                    before_values = {"name":previous["name"],"host":previous["host"],"tcp_port":previous.get("tcpPort",502),"slave":previous.get("slave",1)}
                    after_values = {"name":device_name,"host":host,"tcp_port":tcp_port,"slave":slave}
                    changed = {key:value for key,value in after_values.items() if value != before_values[key]}
                    if changed:
                        record_configuration_audit(
                            db,customer_id,"edge" if forced_site_id else "customer_user","device.updated","device",requested_id,
                            {"name":device_name,"before":{key:before_values[key] for key in changed},"after":changed},
                            None if forced_site_id else str(user["id"]),
                        )
                else:
                    device_id = f"device-{secrets.token_hex(8)}"
                    db.execute(
                        """
                        INSERT INTO devices(
                            id,customer_id,cleanroom_id,site_id,name,host,tcp_port,slave,sort_order
                        ) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)
                        """,
                        (device_id,customer_id,room_id,site_id,device_name,host,tcp_port,slave,sort_order),
                    )
                    record_configuration_audit(
                        db,customer_id,"edge" if forced_site_id else "customer_user","device.created","device",device_id,
                        {"name":device_name,"cleanroom_id":room_id,"site_id":site_id,"host":host,"tcp_port":tcp_port,"slave":slave},
                        None if forced_site_id else str(user["id"]),
                    )
            removed = set(known_devices) - retained
            if removed:
                if forced_site_id:
                    db.execute(
                        """
                        UPDATE devices SET enabled=false,disabled_at=clock_timestamp(),updated_at=clock_timestamp()
                        WHERE customer_id=%s AND cleanroom_id=%s AND site_id=%s
                          AND id = ANY(%s)
                        """,
                        (customer_id, room_id, site_id, list(removed)),
                    )
                else:
                    db.execute(
                        """
                        UPDATE devices SET enabled=false,disabled_at=clock_timestamp(),updated_at=clock_timestamp()
                        WHERE customer_id=%s AND cleanroom_id=%s AND id = ANY(%s)
                        """,
                        (customer_id, room_id, list(removed)),
                    )
                for removed_id in removed:
                    record_configuration_audit(
                        db,customer_id,"edge" if forced_site_id else "customer_user","device.deleted","device",removed_id,
                        {"name":known_devices[removed_id]["name"],"history_retained":True},
                        None if forced_site_id else str(user["id"]),
                    )
            limits = room.get("thresholds") or {}
            try:
                values = normalise_alarm_thresholds(
                    limits,
                    known[room_id].get("thresholds") or {},
                )
            except ValueError as exc:
                raise HTTPException(status_code=400, detail=str(exc)) from exc
            db.execute(
                """
                UPDATE thresholds SET profile_name=%s,
                    particle_0_3_max=%s, particle_0_3_enabled=%s,
                    particle_0_5_max=%s, particle_0_5_enabled=%s,
                    particle_1_0_max=%s, particle_1_0_enabled=%s,
                    particle_2_5_max=%s, particle_2_5_enabled=%s,
                    particle_5_0_max=%s, particle_5_0_enabled=%s,
                    particle_10_0_max=%s, particle_10_0_enabled=%s,
                    particle_5_max=%s,
                    temperature_min=%s, temperature_max=%s, humidity_min=%s,
                    humidity_max=%s, alarm_delay_seconds=%s
                WHERE cleanroom_id=%s AND customer_id=%s
                """,
                (
                    values["profile_name"],
                    values["particle_0_3_max"], values["particle_0_3_enabled"],
                    values["particle_0_5_max"], values["particle_0_5_enabled"],
                    values["particle_1_0_max"], values["particle_1_0_enabled"],
                    values["particle_2_5_max"], values["particle_2_5_enabled"],
                    values["particle_5_0_max"], values["particle_5_0_enabled"],
                    values["particle_10_0_max"], values["particle_10_0_enabled"],
                    values["particle_0_5_max"],
                    values["temperature_min"], values["temperature_max"],
                    values["humidity_min"], values["humidity_max"],
                    values["alarm_delay_seconds"], room_id, customer_id,
                ),
            )
        db.execute(
            "UPDATE sites SET config_version=config_version+1 WHERE customer_id=%s",
            (customer_id,),
        )
    saved = customer_configuration(customer_id)
    if forced_site_id:
        saved = configuration_for_site(saved, forced_site_id)
    return {"ok": True, "data": saved}


@app.post("/api/config")
def update_customer_config(
    payload: CustomerSettings, user: dict[str, object] = Depends(topology_manager)
) -> dict[str, object]:
    return apply_customer_config(payload, user, allow_connection_changes=False)


@app.post("/api/v1/edge/config")
def update_edge_config(
    payload: CustomerSettings, identity: dict[str, str] = Depends(edge_identity)
) -> dict[str, object]:
    """Allow a site collector to update only its own customer's device configuration."""
    return apply_customer_config(
        payload,
        {"customer_id": identity["customer_id"]},
        allow_connection_changes=True,
        forced_site_id=identity["site_id"],
    )


@app.post("/api/discovery/jobs")
def create_discovery_job(
    payload: DiscoveryRequest, user: dict[str, object] = Depends(topology_manager)
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    requested_cidr = validated_discovery_cidr(payload.cidr)
    with connect() as db:
        room = db.execute(
            "SELECT id FROM cleanrooms WHERE id=%s AND customer_id=%s",
            (payload.cleanroom_id, customer_id),
        ).fetchone()
        if room is None:
            raise HTTPException(status_code=404, detail="Cleanroom not found")
        site = db.execute(
            """
            SELECT site_id FROM devices
            WHERE customer_id=%s AND cleanroom_id=%s AND enabled=true AND site_id IS NOT NULL
            ORDER BY sort_order,id LIMIT 1
            """,
            (customer_id, payload.cleanroom_id),
        ).fetchone()
        if site is None:
            site = db.execute(
                "SELECT id FROM sites WHERE customer_id=%s ORDER BY created_at,id LIMIT 1",
                (customer_id,),
            ).fetchone()
        if site is None:
            raise HTTPException(status_code=400, detail="No site collector is configured for this customer")
        site_id = str(site[0])
        db.execute(
            """
            UPDATE device_discovery_jobs
            SET status='failed',completed_at=now(),error='Collector did not complete the scan in time'
            WHERE customer_id=%s AND site_id=%s AND status IN ('pending','running')
              AND requested_at < now() - interval '5 minutes'
            """,
            (customer_id, site_id),
        )
        existing = db.execute(
            f"""
            SELECT {DISCOVERY_JOB_COLUMNS} FROM device_discovery_jobs
            WHERE customer_id=%s AND site_id=%s AND status IN ('pending','running')
            ORDER BY requested_at DESC LIMIT 1
            """,
            (customer_id, site_id),
        ).fetchone()
        if existing is not None:
            return {"ok": True, "data": discovery_job_dict(existing)}
        job_id = uuid4()
        row = db.execute(
            f"""
            INSERT INTO device_discovery_jobs(
                id,customer_id,site_id,cleanroom_id,created_by,requested_cidr,tcp_port
            ) VALUES(%s,%s,%s,%s,%s,%s,%s)
            ON CONFLICT (site_id) WHERE status IN ('pending','running')
            DO NOTHING
            RETURNING {DISCOVERY_JOB_COLUMNS}
            """,
            (
                job_id, customer_id, site_id, payload.cleanroom_id, user["id"],
                requested_cidr, payload.tcp_port,
            ),
        ).fetchone()
        if row is None:
            row = db.execute(
                f"""
                SELECT {DISCOVERY_JOB_COLUMNS} FROM device_discovery_jobs
                WHERE customer_id=%s AND site_id=%s AND status IN ('pending','running')
                ORDER BY requested_at DESC LIMIT 1
                """,
                (customer_id, site_id),
            ).fetchone()
        if row is None:
            raise HTTPException(status_code=409, detail="Discovery request could not be scheduled")
    return {"ok": True, "data": discovery_job_dict(row)}


@app.get("/api/discovery/jobs/latest")
def latest_discovery_job(
    cleanroom_id: str, user: dict[str, object] = Depends(customer_user)
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    with connect() as db:
        row = db.execute(
            f"""
            SELECT {DISCOVERY_JOB_COLUMNS} FROM device_discovery_jobs
            WHERE customer_id=%s AND cleanroom_id=%s
            ORDER BY requested_at DESC LIMIT 1
            """,
            (customer_id, cleanroom_id),
        ).fetchone()
        if row is None:
            return {"ok": True, "data": None}
        result_rows = db.execute(
            """
            SELECT host,tcp_port,slave,verified,modbus_responded,protocol_compatible,
                identity_verified,firmware_raw,particle_unit_code,particle_unit_label,
                unit_supported,latency_ms
            FROM device_discovery_results WHERE job_id=%s
            ORDER BY string_to_array(host,'.')::int[]
            """,
            (row[0],),
        ).fetchall()
    results = [
        {
            "host": item[0], "tcp_port": item[1], "slave": item[2],
            "verified": item[3], "modbus_responded": item[4],
            "protocol_compatible": item[5], "identity_verified": item[6],
            "firmware_raw": item[7], "particle_unit_code": item[8],
            "particle_unit_label": item[9], "unit_supported": item[10],
            "latency_ms": item[11],
        }
        for item in result_rows
    ]
    return {"ok": True, "data": discovery_job_dict(row, results)}


@app.get("/api/v1/edge/discovery/jobs")
def claim_edge_discovery_job(
    identity: dict[str, str] = Depends(edge_identity)
) -> dict[str, object]:
    with connect() as db:
        db.execute(
            """
            DELETE FROM device_discovery_jobs
            WHERE customer_id=%s AND site_id=%s
              AND status IN ('completed','failed')
              AND requested_at < now()-(%s * interval '1 day')
            """,
            (
                identity["customer_id"], identity["site_id"],
                DISCOVERY_HISTORY_RETENTION_DAYS,
            ),
        )
        db.execute(
            """
            UPDATE device_discovery_jobs
            SET status='failed',completed_at=now(),error='Collector did not complete the scan in time'
            WHERE customer_id=%s AND site_id=%s AND status='running'
              AND started_at < now() - interval '5 minutes'
            """,
            (identity["customer_id"], identity["site_id"]),
        )
        active = db.execute(
            """
            SELECT 1 FROM device_discovery_jobs
            WHERE customer_id=%s AND site_id=%s AND status IN ('pending','running')
            LIMIT 1
            """,
            (identity["customer_id"], identity["site_id"]),
        ).fetchone()
        recent_automatic = db.execute(
            """
            SELECT 1 FROM device_discovery_jobs
            WHERE customer_id=%s AND site_id=%s AND source='automatic'
              AND requested_at > now()-(%s * interval '1 second')
            LIMIT 1
            """,
            (
                identity["customer_id"], identity["site_id"],
                AUTO_DISCOVERY_INTERVAL_SECONDS,
            ),
        ).fetchone()
        if active is None and recent_automatic is None:
            db.execute(
                """
                INSERT INTO device_discovery_jobs(
                    id,customer_id,site_id,cleanroom_id,created_by,
                    requested_cidr,tcp_port,source
                ) VALUES(%s,%s,%s,NULL,NULL,NULL,502,'automatic')
                ON CONFLICT (site_id) WHERE status IN ('pending','running')
                DO NOTHING
                """,
                (uuid4(), identity["customer_id"], identity["site_id"]),
            )
        row = db.execute(
            f"""
            WITH next_job AS (
                SELECT id FROM device_discovery_jobs
                WHERE customer_id=%s AND site_id=%s AND status='pending'
                ORDER BY requested_at FOR UPDATE SKIP LOCKED LIMIT 1
            )
            UPDATE device_discovery_jobs AS job
            SET status='running',started_at=now(),error=NULL
            FROM next_job WHERE job.id=next_job.id
            RETURNING {','.join(f'job.{column.strip()}' for column in DISCOVERY_JOB_COLUMNS.split(','))}
            """,
            (identity["customer_id"], identity["site_id"]),
        ).fetchone()
    if row is None:
        return {"ok": True, "data": None}
    return {
        "ok": True,
        "data": {
            "id": str(row[0]), "requested_cidr": row[2], "tcp_port": row[4],
            "cleanroom_id": row[10], "site_id": row[11], "source": row[13],
        },
    }


@app.post("/api/v1/edge/discovery/jobs/{job_id}/result")
def complete_edge_discovery_job(
    job_id: UUID, payload: DiscoveryCompletion,
    identity: dict[str, str] = Depends(edge_identity),
) -> dict[str, object]:
    raw_scanned_cidrs = list(payload.scanned_cidrs)
    if payload.scanned_cidr and payload.scanned_cidr not in raw_scanned_cidrs:
        raw_scanned_cidrs.insert(0, payload.scanned_cidr)
    scanned_cidrs = (
        [validated_discovery_cidr(value) for value in raw_scanned_cidrs]
        if payload.status == "completed" else []
    )
    scanned_cidrs = [value for value in scanned_cidrs if value is not None]
    if payload.status == "completed" and not scanned_cidrs:
        raise HTTPException(status_code=400, detail="Completed scans must include a scanned network")
    networks = [ipaddress.ip_network(value) for value in scanned_cidrs]
    validated_results: list[tuple[object, ...]] = []
    for item in payload.results:
        try:
            address = ipaddress.ip_address(item.host)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=f"Invalid discovered IP address: {item.host}") from exc
        if address.version != 4 or not any(address in allowed for allowed in PRIVATE_DEVICE_NETWORKS):
            raise HTTPException(status_code=400, detail="Discovered devices must use private IPv4 addresses")
        if networks and not any(address in network for network in networks):
            raise HTTPException(status_code=400, detail="Discovery result is outside the scanned networks")
        validated_results.append(
            (
                str(address), item.tcp_port, item.slave, item.verified,
                item.modbus_responded, item.protocol_compatible, item.identity_verified,
                item.firmware_raw, item.particle_unit_code, item.particle_unit_label,
                item.unit_supported, item.latency_ms,
            )
        )
    if payload.status == "failed" and not (payload.error or "").strip():
        raise HTTPException(status_code=400, detail="Failed scans must include an error")
    with connect() as db:
        job = db.execute(
            """
            SELECT tcp_port,status FROM device_discovery_jobs
            WHERE id=%s AND customer_id=%s AND site_id=%s FOR UPDATE
            """,
            (job_id, identity["customer_id"], identity["site_id"]),
        ).fetchone()
        if job is None:
            raise HTTPException(status_code=404, detail="Discovery job not found")
        if any(item[1] != int(job[0]) for item in validated_results):
            raise HTTPException(status_code=400, detail="Discovery result port does not match the requested port")
        db.execute("DELETE FROM device_discovery_results WHERE job_id=%s", (job_id,))
        if validated_results:
            with db.cursor() as cursor:
                cursor.executemany(
                    """
                    INSERT INTO device_discovery_results(
                        job_id,host,tcp_port,slave,verified,modbus_responded,
                        protocol_compatible,identity_verified,firmware_raw,
                        particle_unit_code,particle_unit_label,unit_supported,latency_ms
                    ) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                    """,
                    [(job_id, *item) for item in validated_results],
                )
        db.execute(
            """
            UPDATE device_discovery_jobs
            SET status=%s,scanned_cidr=%s,scanned_cidrs=%s::jsonb,
                completed_at=now(),result_count=%s,error=%s
            WHERE id=%s
            """,
            (
                payload.status, scanned_cidrs[0] if scanned_cidrs else None,
                json.dumps(scanned_cidrs), len(validated_results),
                (payload.error or "").strip()[:500] or None, job_id,
            ),
        )
    return {"ok": True, "accepted": len(validated_results)}


def edge_configuration_payload(identity: dict[str, str]) -> dict[str, object]:
    with connect() as db:
        # List and revision must describe the same committed snapshot. Otherwise
        # a collector may permanently cache old devices under the new revision.
        db.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
        rooms = configuration_for_site(
            customer_configuration(identity["customer_id"], db=db), identity["site_id"],
        )
        row = db.execute(
            "SELECT config_version,name FROM sites WHERE id=%s AND customer_id=%s",
            (identity["site_id"], identity["customer_id"]),
        ).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Site not found")
    return {
        "customer_id": identity["customer_id"], "site_id": identity["site_id"],
        "site_name": str(row[1]), "revision": int(row[0]), "rooms": rooms,
    }


@app.post("/api/v1/edge/cleanrooms", status_code=201)
def edge_create_cleanroom(
    payload: CleanroomCreate,
    identity: dict[str, str] = Depends(edge_identity),
) -> dict[str, object]:
    with connect() as db:
        site_count = int(db.execute(
            "SELECT COUNT(*) FROM sites WHERE customer_id=%s", (identity["customer_id"],)
        ).fetchone()[0])
    if site_count != 1:
        raise HTTPException(
            status_code=409,
            detail="Create cleanrooms in the cloud administrator console for multi-site customers",
        )
    create_cloud_cleanroom(
        identity["customer_id"], payload.name, actor_kind="edge",
    )
    return {"ok": True, "data": edge_configuration_payload(identity)}


@app.post("/api/v1/edge/devices", status_code=201)
def edge_create_device(
    payload: DeviceCreate,
    identity: dict[str, str] = Depends(edge_identity),
) -> dict[str, object]:
    create_cloud_device(
        identity["customer_id"],
        {
            "cleanroom_id": payload.cleanroom_id, "name": payload.name,
            "host": payload.host, "tcp_port": payload.tcp_port, "slave": payload.slave,
        },
        actor_kind="edge", forced_site_id=identity["site_id"],
    )
    return {"ok": True, "data": edge_configuration_payload(identity)}


@app.get("/api/v1/edge/config")
def edge_config(identity: dict[str, str] = Depends(edge_identity)) -> dict[str, object]:
    return {"ok": True, "data": edge_configuration_payload(identity)}


@app.post("/api/v1/edge/heartbeat")
def edge_heartbeat(
    payload: CollectorHeartbeat,
    identity: dict[str, str] = Depends(edge_identity),
) -> dict[str, object]:
    if payload.site_id != identity["site_id"]:
        raise HTTPException(status_code=403, detail="Heartbeat site does not match edge identity")
    if payload.device_online > payload.device_total:
        raise HTTPException(status_code=422, detail="Online device count exceeds total devices")
    with connect() as db:
        row = db.execute(
            "SELECT config_version FROM sites WHERE id=%s AND customer_id=%s FOR UPDATE",
            (identity["site_id"], identity["customer_id"]),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Collector node not found")
        desired_revision = int(row[0])
        if (
            payload.applied_config_version is not None
            and payload.applied_config_version > desired_revision
        ):
            raise HTTPException(
                status_code=409,
                detail="Applied configuration revision is newer than the cloud revision",
            )
        if payload.config_apply_status == "applied" and payload.applied_config_version is None:
            raise HTTPException(
                status_code=422,
                detail="Applied configuration status requires an applied revision",
            )
        status = payload.model_dump()
        status["instance_id"] = identity["instance_id"]
        status["clock_offset_seconds"] = round(time.time() - payload.collector_time, 3)
        db.execute(
            """
            UPDATE sites SET
              collector_status=%s::jsonb,last_heartbeat_at=now(),
              applied_config_version=%s,config_apply_status=%s,
              config_apply_error=%s,updated_at=now()
            WHERE id=%s AND customer_id=%s
            """,
            (
                json.dumps(status, ensure_ascii=False), payload.applied_config_version,
                payload.config_apply_status,
                payload.config_apply_error if payload.config_apply_status == "failed" else None,
                identity["site_id"], identity["customer_id"],
            ),
        )
    return {
        "ok": True,
        "data": {
            "desired_config_version": desired_revision,
            "heartbeat_at": time.time(),
        },
    }


READING_COLUMNS = """
record_uuid, cleanroom_id, cleanroom_name, device_id, device_name, measured_at,
source, particles, environment, alarm_status, alarm_details, cleanliness_code, cleanliness_label,
particle_unit_code, particle_unit_label, protocol_profile
"""


@app.get("/api/latest")
def customer_latest(
    cleanroom_id: str | None = None,
    user: dict[str, object] = Depends(customer_user),
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    clauses = ["latest.customer_id=%s"]
    params: list[object] = [customer_id]
    if cleanroom_id:
        clauses.append("latest.cleanroom_id=%s")
        params.append(cleanroom_id)
    with connect() as db:
        rows = db.execute(
            f"""
            SELECT NULL::uuid, latest.cleanroom_id, latest.cleanroom_name,
                latest.device_id, latest.device_name, latest.measured_at, latest.source,
                latest.particles, latest.environment, latest.alarm_status,
                latest.alarm_details, latest.cleanliness_code, latest.cleanliness_label,
                latest.particle_unit_code, latest.particle_unit_label, latest.protocol_profile
            FROM latest_readings AS latest
            JOIN devices AS device
              ON device.customer_id=latest.customer_id AND device.id=latest.device_id
             AND device.enabled=true
            WHERE {' AND '.join(clauses)}
            ORDER BY latest.cleanroom_id, latest.device_id
            """,
            tuple(params),
        ).fetchall()
    return {"ok": True, "data": annotate_maintenance([reading_dict(row) for row in rows], customer_id, current=True)}


@app.get("/api/history")
def customer_history(
    cleanroom: str | None = None,
    cleanroom_ids: str | None = None,
    device_ids: str | None = None,
    start: float | None = None,
    end: float | None = None,
    limit: int = Query(default=5000, ge=1, le=5000),
    user: dict[str, object] = Depends(customer_user),
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    clauses = ["customer_id=%s"]
    params: list[object] = [customer_id]
    if cleanroom:
        clauses.append("cleanroom_id=(SELECT id FROM cleanrooms WHERE customer_id=%s AND name=%s LIMIT 1)")
        params.extend((customer_id, cleanroom))
    requested_room_ids = parse_scope_parameter(cleanroom_ids)
    requested_device_ids = parse_scope_parameter(device_ids)
    if requested_room_ids:
        clauses.append(f"cleanroom_id IN ({','.join('%s' for _ in requested_room_ids)})")
        params.extend(requested_room_ids)
    if requested_device_ids:
        clauses.append(f"device_id IN ({','.join('%s' for _ in requested_device_ids)})")
        params.extend(requested_device_ids)
    if start is not None:
        clauses.append("measured_at >= to_timestamp(%s)")
        params.append(start)
    if end is not None:
        clauses.append("measured_at <= to_timestamp(%s)")
        params.append(end)
    params.append(limit)
    with connect() as db:
        rows = db.execute(
            f"SELECT {READING_COLUMNS} FROM readings WHERE {' AND '.join(clauses)} ORDER BY measured_at DESC LIMIT %s",
            tuple(params),
        ).fetchall()
    return {"ok": True, "data": annotate_maintenance([reading_dict(row) for row in rows], customer_id, current=False)}


@app.get("/api/trends")
def customer_trends(
    cleanroom_ids: str | None = None,
    device_ids: str | None = None,
    start: float | None = None,
    end: float | None = None,
    max_points: int = Query(default=240, ge=20, le=600),
    user: dict[str, object] = Depends(customer_user),
) -> dict[str, object]:
    """Return an evenly sampled time series for every selected device."""
    customer_id = str(user["customer_id"])
    end = end if end is not None else time.time()
    start = start if start is not None else end - 3600
    if start >= end:
        raise HTTPException(status_code=400, detail="Start time must be before end time")

    clauses = ["customer_id=%s", "measured_at>=to_timestamp(%s)", "measured_at<=to_timestamp(%s)"]
    params: list[object] = [customer_id, start, end]
    requested_room_ids = parse_scope_parameter(cleanroom_ids)
    requested_device_ids = parse_scope_parameter(device_ids)
    if requested_room_ids:
        clauses.append(f"cleanroom_id IN ({','.join('%s' for _ in requested_room_ids)})")
        params.extend(requested_room_ids)
    if requested_device_ids:
        clauses.append(f"device_id IN ({','.join('%s' for _ in requested_device_ids)})")
        params.extend(requested_device_ids)
    params.append(max_points)

    with connect() as db:
        rows = db.execute(
            f"""
            WITH filtered AS (
                SELECT {READING_COLUMNS}, COALESCE(NULLIF(device_id, ''), device_name) AS sample_device_key
                FROM readings
                WHERE {' AND '.join(clauses)}
            ),
            sequenced AS (
                SELECT filtered.*,
                    ROW_NUMBER() OVER (
                        PARTITION BY sample_device_key ORDER BY measured_at ASC
                    ) AS sample_sequence,
                    COUNT(*) OVER (PARTITION BY sample_device_key) AS sample_count
                FROM filtered
            ),
            bucketed AS (
                SELECT sequenced.*,
                    CASE WHEN sample_count <= 1 THEN 0 ELSE FLOOR(
                        ((sample_sequence - 1) * (%s - 1))::numeric / (sample_count - 1)
                    )::integer END AS sample_bucket
                FROM sequenced
            ),
            ranked AS (
                SELECT bucketed.*,
                    ROW_NUMBER() OVER (
                        PARTITION BY sample_device_key, sample_bucket ORDER BY measured_at ASC
                    ) AS bucket_rank
                FROM bucketed
            )
            SELECT {READING_COLUMNS}
            FROM ranked
            WHERE bucket_rank=1
            ORDER BY measured_at ASC
            """,
            tuple(params),
        ).fetchall()
    return {"ok": True, "data": annotate_maintenance([reading_dict(row) for row in rows], customer_id, current=False)}


@app.get("/api/alarms")
def customer_alarms(
    cleanroom_id: str | None = None, limit: int = Query(default=200, ge=1, le=2000),
    user: dict[str, object] = Depends(customer_user),
) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    clauses, params = ["customer_id=%s"], [customer_id]
    if cleanroom_id:
        clauses.append("cleanroom_id=%s")
        params.append(cleanroom_id)
    params.append(limit)
    with connect() as db:
        rows = db.execute(
            f"""
            SELECT event_uuid, cleanroom_id, device_id, metric, started_at, ended_at,
                trigger_value, peak_value, limit_description
            FROM alarm_events WHERE {' AND '.join(clauses)}
            ORDER BY started_at DESC LIMIT %s
            """,
            tuple(params),
        ).fetchall()
    data = [
        {
            "event_uuid": str(row[0]), "cleanroom_id": row[1], "device_id": row[2],
            "metric": row[3], "started_at": row[4].timestamp(),
            "ended_at": row[5].timestamp() if row[5] else None,
            "trigger_value": row[6], "peak_value": row[7], "limit_description": row[8],
        }
        for row in rows
    ]
    return {"ok": True, "data": annotate_maintenance(data, customer_id)}


@app.get("/api/stats")
def customer_stats(user: dict[str, object] = Depends(customer_user)) -> dict[str, object]:
    customer_id = str(user["customer_id"])
    with connect() as db:
        row = db.execute(
            """
            SELECT COUNT(*), MAX(measured_at),
                (SELECT COUNT(*) FROM alarm_events WHERE customer_id=%s)
            FROM readings WHERE customer_id=%s
            """,
            (customer_id, customer_id),
        ).fetchone()
    return {
        "ok": True,
        "data": {
            "readings": row[0], "latest_reading": row[1].timestamp() if row[1] else None,
            "alarms": row[2], "pending_sync": 0, "storage_mode": "Cloud PostgreSQL",
        },
    }


@app.get("/api/history/export-xlsx")
def customer_excel_export(
    start: float | None = None,
    end: float | None = None,
    cleanroom_ids: str | None = None,
    device_ids: str | None = None,
    locale: str = Query(default="en-US", max_length=16),
    user: dict[str, object] = Depends(customer_user),
) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    customer_id = str(user["customer_id"])
    locale = normalize_report_locale(locale)
    catalog = report_catalog(locale)
    requested_room_ids = set(parse_scope_parameter(cleanroom_ids))
    requested_device_ids = parse_scope_parameter(device_ids)
    requested_device_id_set = set(requested_device_ids)
    config = customer_configuration(customer_id, include_disabled=True)
    if requested_room_ids:
        config = [room for room in config if str(room["id"]) in requested_room_ids]
    if requested_device_id_set:
        config = [
            room for room in config
            if any(str(device["id"]) in requested_device_id_set for device in room.get("devices", []))
        ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary_sheet_name = str(catalog["summary_sheet"])
    summary = workbook.create_sheet(summary_sheet_name)
    summary_headers = list(catalog["summary_headers"])
    summary.append(summary_headers)
    for cell in summary[1]:
        cell.font = Font(bold=True)
    particle_unit = report_particle_unit(locale)
    metric_specs = [
        ("0.3 µm", particle_unit, "particles", "pm_0_3_um"),
        ("0.5 µm", particle_unit, "particles", "pm_0_5_um"),
        ("1.0 µm", particle_unit, "particles", "pm_1_0_um"),
        ("2.5 µm", particle_unit, "particles", "pm_2_5_um"),
        ("5.0 µm", particle_unit, "particles", "pm_5_0_um"),
        ("10.0 µm", particle_unit, "particles", "pm_10_0_um"),
        (report_metric_name("temperature", locale), "°C", "environment", "temperature"),
        (report_metric_name("humidity", locale), "%RH", "environment", "humidity"),
    ]
    red_fill = PatternFill("solid", fgColor="FDECEB")
    used_titles = {summary_sheet_name}
    for room in config:
        base_title = re.sub(r"[\\/*?:\[\]]", "_", str(room["name"])).strip()[:31]
        base_title = base_title or str(catalog["cleanroom_fallback"])
        title = base_title
        suffix = 2
        while title in used_titles:
            marker = f" {suffix}"
            title = f"{base_title[:31-len(marker)]}{marker}"
            suffix += 1
        used_titles.add(title)
        sheet = workbook.create_sheet(title or str(catalog["cleanroom_fallback"]))
        headers = list(catalog["table_headers"])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        clauses = ["customer_id=%s", "cleanroom_id=%s"]
        params: list[object] = [customer_id, room["id"]]
        if start is not None:
            clauses.append("measured_at>=to_timestamp(%s)")
            params.append(start)
        if end is not None:
            clauses.append("measured_at<=to_timestamp(%s)")
            params.append(end)
        if requested_device_ids:
            clauses.append(f"device_id IN ({','.join('%s' for _ in requested_device_ids)})")
            params.extend(requested_device_ids)
        with connect() as db:
            record_count = db.execute(
                f"SELECT COUNT(*) FROM readings WHERE {' AND '.join(clauses)}", tuple(params)
            ).fetchone()[0]
            if record_count > 100000:
                raise HTTPException(
                    status_code=413,
                    detail=f"{room['name']} has more than 100,000 rows; select a shorter time range",
                )
            rows = db.execute(
                f"""
                SELECT device_id,device_name,measured_at,particles,environment,
                    alarm_status,alarm_details FROM readings
                WHERE {' AND '.join(clauses)} ORDER BY measured_at
                """,
                tuple(params),
            ).fetchall()
            alarm_rows = db.execute(
                f"""
                SELECT device_id,COUNT(*) FROM alarm_events
                WHERE {' AND '.join(clauses).replace('measured_at', 'started_at')}
                GROUP BY device_id
                """,
                tuple(params),
            ).fetchall()
        alarm_counts = dict(alarm_rows)
        display_names = {device["id"]: device["name"] for device in room["devices"]}
        maintenance_records = annotate_maintenance([{"device_id":r[0],"timestamp":r[2].timestamp()} for r in rows],customer_id)
        by_device: dict[str, list[tuple[Any, ...]]] = {}
        for row, marked in zip(rows, maintenance_records):
            by_device.setdefault(row[0], []).append(row)
            details = report_alarm_details(row[6] or [], locale)
            if marked.get("maintenance"):
                label = "维护期间" if locale == "zh-CN" else "During maintenance"
                details = f"{details} · {label}: {marked['maintenance']['reason']}".strip(" ·")
            particles, environment = row[3] or {}, row[4] or {}
            sheet.append([
                row[2].replace(tzinfo=None), row[1],
                particles.get("pm_0_3_um"), particles.get("pm_0_5_um"),
                particles.get("pm_1_0_um"), particles.get("pm_2_5_um"),
                particles.get("pm_5_0_um"), particles.get("pm_10_0_um"),
                environment.get("temperature"), environment.get("humidity"),
                report_status(row[5], locale), details,
            ])
            if row[5] == "ALARM_ACTIVE":
                for cell in sheet[sheet.max_row]:
                    cell.fill = red_fill
        for device_id, device_rows in by_device.items():
            for metric, unit, group, key in metric_specs:
                values = [float(row[3 if group == "particles" else 4][key]) for row in device_rows if key in (row[3 if group == "particles" else 4] or {})]
                if values:
                    summary.append([
                        room["name"], display_names.get(device_id, device_id), metric, unit,
                        sum(values) / len(values), min(values), max(values), alarm_counts.get(device_id, 0),
                        len(values), device_rows[0][2].replace(tzinfo=None),
                        device_rows[-1][2].replace(tzinfo=None),
                    ])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    summary.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=hawkhive-cleanroom-report-{locale}.xlsx",
            "Content-Language": locale,
        },
    )


@app.post("/api/v1/readings/batch")
def ingest_batch(batch: ReadingBatch, identity: dict[str, str] = Depends(edge_identity)) -> dict[str, object]:
    if batch.site_id != identity["site_id"]:
        raise HTTPException(status_code=403, detail="Token is not valid for this site")
    accepted: list[str] = []
    with connect() as db:
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR KEY SHARE", (identity["customer_id"],))
        for item in batch.readings:
            if item.site_id != identity["site_id"] or item.customer_id != identity["customer_id"]:
                raise HTTPException(status_code=403, detail="Reading tenant does not match edge token")
            cleanroom_name, device_name = authorize_ingest_target(
                db, identity, item.cleanroom_id, item.device_id, historical_at=item.measured_at,
            )
            db.execute(
                """
                INSERT INTO readings(
                    record_uuid, customer_id, site_id, cleanroom_id, cleanroom_name,
                    device_id, device_name, measured_at, source, particles, environment,
                    alarm_status, alarm_details, cleanliness_code, cleanliness_label,
                    particle_unit_code, particle_unit_label, protocol_profile
                ) VALUES (
                    %s::uuid, %s, %s, %s, %s, %s, %s, to_timestamp(%s), %s,
                    %s::jsonb, %s::jsonb, %s, %s::jsonb, %s, %s, %s, %s, %s
                ) ON CONFLICT(record_uuid) DO NOTHING
                """,
                (
                    item.record_uuid, identity["customer_id"], identity["site_id"],
                    item.cleanroom_id, cleanroom_name, item.device_id, device_name,
                    item.measured_at, item.source, json.dumps(item.particles),
                    json.dumps(item.environment), item.alarm_status,
                    json.dumps(item.alarm_details), item.cleanliness_code,
                    item.cleanliness_label, item.particle_unit_code,
                    item.particle_unit_label, item.protocol_profile,
                ),
            )
            accepted.append(str(item.record_uuid))
    return {"ok": True, "accepted": accepted}


@app.post("/api/v1/latest/batch")
def ingest_latest(batch: LatestBatch, identity: dict[str, str] = Depends(edge_identity)) -> dict[str, object]:
    if batch.site_id != identity["site_id"]:
        raise HTTPException(status_code=403, detail="Token is not valid for this site")
    ignored_device_ids: list[str] = []
    with connect() as db:
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR KEY SHARE", (identity["customer_id"],))
        for item in batch.readings:
            if item.site_id != identity["site_id"] or item.customer_id != identity["customer_id"]:
                raise HTTPException(status_code=403, detail="Reading tenant does not match edge token")
            retired = db.execute(
                """SELECT 1 FROM device_assignment_periods p
                   WHERE p.customer_id=%s AND p.device_id=%s AND p.site_id=%s AND p.cleanroom_id=%s
                     AND NOT EXISTS (SELECT 1 FROM device_assignment_periods current
                         WHERE current.device_id=p.device_id AND current.site_id=p.site_id
                           AND current.cleanroom_id=p.cleanroom_id AND current.ended_at IS NULL
                           AND current.started_at<=to_timestamp(%s)) LIMIT 1""",
                (identity["customer_id"],item.device_id,identity["site_id"],item.cleanroom_id,item.measured_at),
            ).fetchone()
            if retired:
                # Old collectors may finish an in-flight poll after applying the
                # removal. Acknowledge and discard only an owned retired target;
                # one stale cache must not block healthy devices in this batch.
                ignored_device_ids.append(item.device_id)
                continue
            cleanroom_name, device_name = authorize_ingest_target(
                db, identity, item.cleanroom_id, item.device_id,
            )
            db.execute(
                """
                INSERT INTO latest_readings(
                    customer_id, site_id, cleanroom_id, cleanroom_name, device_id,
                    device_name, measured_at, source, particles, environment,
                    alarm_status, alarm_details, cleanliness_code, cleanliness_label,
                    particle_unit_code, particle_unit_label, protocol_profile
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, to_timestamp(%s), %s,
                    %s::jsonb, %s::jsonb, %s, %s::jsonb, %s, %s, %s, %s, %s
                )
                ON CONFLICT(customer_id, site_id, device_id) DO UPDATE SET
                    cleanroom_id=excluded.cleanroom_id, cleanroom_name=excluded.cleanroom_name,
                    device_name=excluded.device_name, measured_at=excluded.measured_at,
                    received_at=now(), source=excluded.source, particles=excluded.particles,
                    environment=excluded.environment, alarm_status=excluded.alarm_status,
                    alarm_details=excluded.alarm_details,
                    cleanliness_code=excluded.cleanliness_code,
                    cleanliness_label=excluded.cleanliness_label,
                    particle_unit_code=excluded.particle_unit_code,
                    particle_unit_label=excluded.particle_unit_label,
                    protocol_profile=excluded.protocol_profile
                WHERE excluded.measured_at >= latest_readings.measured_at
                """,
                (
                    identity["customer_id"], identity["site_id"], item.cleanroom_id,
                    cleanroom_name, item.device_id, device_name, item.measured_at,
                    item.source, json.dumps(item.particles), json.dumps(item.environment),
                    item.alarm_status, json.dumps(item.alarm_details), item.cleanliness_code,
                    item.cleanliness_label, item.particle_unit_code,
                    item.particle_unit_label, item.protocol_profile,
                ),
            )
    return {"ok": True, "accepted": len(batch.readings), "ignored_device_ids": ignored_device_ids}


@app.post("/api/v1/alarms/batch")
def ingest_alarms(batch: AlarmBatch, identity: dict[str, str] = Depends(edge_identity)) -> dict[str, object]:
    if batch.site_id != identity["site_id"]:
        raise HTTPException(status_code=403, detail="Token is not valid for this site")
    accepted: list[str] = []
    with connect() as db:
        db.execute("SELECT 1 FROM customers WHERE id=%s FOR KEY SHARE", (identity["customer_id"],))
        for item in batch.events:
            if item.site_id != identity["site_id"] or item.customer_id != identity["customer_id"]:
                raise HTTPException(status_code=403, detail="Alarm tenant does not match edge token")
            authorize_ingest_target(
                db, identity, item.cleanroom_id, item.device_id,
                historical_at=item.ended_at or item.started_at,
            )
            existing = db.execute(
                """SELECT customer_id,site_id,cleanroom_id,device_id
                   FROM alarm_events WHERE event_uuid=%s::uuid""",
                (item.event_uuid,),
            ).fetchone()
            expected_owner = (
                identity["customer_id"], identity["site_id"],
                item.cleanroom_id, item.device_id,
            )
            if existing is not None and tuple(str(value) for value in existing) != expected_owner:
                raise HTTPException(status_code=409, detail="Alarm event id belongs to another target")
            alarm_write = db.execute(
                """
                INSERT INTO alarm_events(
                    event_uuid,customer_id,site_id,cleanroom_id,device_id,source,metric,
                    started_at,ended_at,trigger_value,peak_value,limit_description
                ) VALUES(
                    %s::uuid,%s,%s,%s,%s,%s,%s,to_timestamp(%s),
                    to_timestamp(%s::double precision),%s,%s,%s
                )
                ON CONFLICT(event_uuid) DO UPDATE SET ended_at=COALESCE(alarm_events.ended_at,excluded.ended_at),
                    peak_value=excluded.peak_value,limit_description=excluded.limit_description,
                    received_at=now()
                WHERE alarm_events.customer_id=excluded.customer_id
                  AND alarm_events.site_id=excluded.site_id
                  AND alarm_events.cleanroom_id=excluded.cleanroom_id
                  AND alarm_events.device_id=excluded.device_id
                """,
                (
                    item.event_uuid,identity["customer_id"],identity["site_id"],
                    item.cleanroom_id,item.device_id,item.source,item.metric,item.started_at,
                    item.ended_at,item.trigger_value,item.peak_value,
                    item.limit_description,
                ),
            )
            if alarm_write.rowcount == 0:
                raise HTTPException(status_code=409, detail="Alarm event id belongs to another target")
            enqueue_alarm(db, identity["customer_id"], item.event_uuid)
            accepted.append(str(item.event_uuid))
    return {"ok": True, "accepted": accepted}


register_email_routes(app, connect, topology_manager, record_configuration_audit)
register_diagnostic_routes(app, connect, topology_manager, edge_identity)
from device_lifecycle import register_device_lifecycle
import sys
register_device_lifecycle(sys.modules[__name__])
from data_management import register_data_management
register_data_management(sys.modules[__name__])

register_update_routes(app)
PUBLIC_DIR = Path(__file__).with_name("public")
if PUBLIC_DIR.exists():
    app.mount("/", StaticFiles(directory=PUBLIC_DIR, html=True), name="customer-dashboard")
