"""Opt-in deletion acceptance against a disposable PostgreSQL database."""
from __future__ import annotations

import os
import io
import threading
from concurrent.futures import ThreadPoolExecutor
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from fastapi.testclient import TestClient

import cloud_api
import dashboard_server
from auth_service import AuthService
from monitoring_service import MonitoringService


def seed_customer(label: str) -> dict:
    customer_id = f"deleteqa-{uuid4().hex}"
    user_id = str(uuid4())
    salt, digest = AuthService.password_hash("local-delete-qa")
    with cloud_api.connect() as db:
        db.execute("INSERT INTO customers(id,name) VALUES(%s,%s)", (customer_id, label))
        db.execute(
            """INSERT INTO customer_users(id,customer_id,username,display_name,password_salt,password_hash)
               VALUES(%s,%s,%s,%s,%s,%s)""",
            (user_id, customer_id, customer_id, label, salt, digest),
        )
        for suffix in ("primary", "other"):
            db.execute("INSERT INTO sites(id,customer_id,name) VALUES(%s,%s,%s)",
                       (f"{customer_id}-{suffix}", customer_id, suffix))
    user = {"id": user_id, "customer_id": customer_id, "role": "customer_admin"}
    room_id = cloud_api.create_cloud_cleanroom(customer_id, "删除验收车间", actor_kind="customer_user", actor_user_id=user_id)
    payload = {"cleanroom_id": room_id, "site_id": f"{customer_id}-primary",
               "name": "Deletion QA device", "host": "192.168.50.30", "tcp_port": 502, "slave": 1}
    device_id = cloud_api.create_cloud_device(customer_id, payload, actor_kind="customer_user", actor_user_id=user_id)
    return {"user": user, "payload": payload, "device_id": device_id}


@unittest.skipUnless(os.environ.get("DCP_DEVICE_DELETE_TEST_DATABASE_URL"), "Opt-in isolated PostgreSQL required")
class DeviceDeletionPostgresTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        dsn = os.environ["DCP_DEVICE_DELETE_TEST_DATABASE_URL"]
        if not dsn.endswith("/devicedeleteqa"):
            raise RuntimeError("Use an isolated database named devicedeleteqa")
        cls.environment = patch.dict(os.environ, {"DATABASE_URL": dsn, "DASHBOARD_SECURE_COOKIE": "0"})
        cls.environment.start()
        cloud_api.initialize_schema()

    @classmethod
    def tearDownClass(cls):
        cls.environment.stop()

    def setUp(self):
        self.fixture = seed_customer(self.id())
        self.user = self.fixture["user"]
        self.device_id = self.fixture["device_id"]
        self.payload = self.fixture["payload"]
        self.client = TestClient(cloud_api.app, raise_server_exceptions=False)
        result = self.client.post("/api/login", json={"username": self.user["customer_id"], "password": "local-delete-qa"})
        self.assertEqual(result.status_code, 200)
        self.url = f"/api/admin/devices/{self.device_id}"

    def tearDown(self):
        self.client.close()

    def versions(self):
        with cloud_api.connect() as db:
            return dict(db.execute("SELECT id,config_version FROM sites WHERE customer_id=%s", (self.user["customer_id"],)).fetchall())

    def test_delete_retains_history_updates_only_owner_and_is_idempotent(self):
        identity = {"customer_id": self.user["customer_id"], "site_id": self.payload["site_id"]}
        record = dict(customer_id=identity["customer_id"], site_id=identity["site_id"],
                      cleanroom_id=self.payload["cleanroom_id"], cleanroom_name="删除验收车间",
                      device_id=self.device_id, device_name=self.payload["name"], measured_at=time.time(),
                      source="device", particles={}, environment={}, alarm_status="normal")
        cloud_api.ingest_batch(cloud_api.ReadingBatch(site_id=identity["site_id"], readings=[dict(record, record_uuid=uuid4())]), identity)
        cloud_api.ingest_latest(cloud_api.LatestBatch(site_id=identity["site_id"], readings=[record]), identity)
        cloud_api.ingest_alarms(cloud_api.AlarmBatch(site_id=identity["site_id"], events=[dict(
            event_uuid=uuid4(), customer_id=identity["customer_id"], site_id=identity["site_id"],
            cleanroom_id=self.payload["cleanroom_id"], device_id=self.device_id, source="device",
            metric="temperature", started_at=time.time(), limit_description="QA",
        )]), identity)
        before = self.versions()
        self.assertEqual(len(self.client.get("/api/latest").json()["data"]), 1)
        response = self.client.delete(self.url)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["data"][0]["devices"], [])
        self.assertEqual(self.client.get("/api/latest").json()["data"], [])
        self.assertEqual(len(self.client.get("/api/history").json()["data"]), 1)
        self.assertEqual(len(self.client.get("/api/alarms").json()["data"]), 1)
        historical_config = self.client.get("/api/config?include_disabled=true").json()["data"]
        self.assertEqual(historical_config[0]["devices"][0]["id"], self.device_id)
        self.assertFalse(historical_config[0]["devices"][0]["enabled"])
        from openpyxl import load_workbook
        export = self.client.get("/api/history/export-xlsx", params={"device_ids": self.device_id})
        self.assertEqual(export.status_code, 200, export.text if export.status_code != 200 else "")
        workbook = load_workbook(io.BytesIO(export.content))
        self.assertEqual(len(workbook.sheetnames), 2)
        self.assertGreater(workbook.worksheets[1].max_row, 1)
        after = self.versions()
        self.assertEqual(after[identity["site_id"]], before[identity["site_id"]] + 1)
        self.assertEqual(after[f'{identity["customer_id"]}-other'], before[f'{identity["customer_id"]}-other'])
        self.assertEqual(self.client.delete(self.url).status_code, 200)
        self.assertEqual(self.versions(), after)
        with cloud_api.connect() as db:
            self.assertEqual(db.execute("SELECT enabled,disabled_at IS NOT NULL FROM devices WHERE id=%s", (self.device_id,)).fetchone(), (False, True))
            audit = db.execute("SELECT actor_user_id,details FROM configuration_audit_events WHERE target_id=%s AND action='device.deleted'", (self.device_id,)).fetchall()
            self.assertEqual(len(audit), 1)
            self.assertEqual(str(audit[0][0]), self.user["id"])
            self.assertTrue(audit[0][1]["history_retained"])
        # Retired live caches are acknowledged without writing, for old collectors.
        retired_live = cloud_api.ingest_latest(cloud_api.LatestBatch(site_id=identity["site_id"], readings=[record]), identity)
        self.assertEqual(retired_live["ignored_device_ids"], [self.device_id])
        self.assertEqual(self.client.get("/api/latest").json()["data"], [])
        added = self.client.post("/api/admin/devices", json=self.payload)
        self.assertEqual(added.status_code, 201, added.text)
        self.assertNotEqual(added.json()["data"][0]["devices"][0]["id"], self.device_id)
        new_id = added.json()["data"][0]["devices"][0]["id"]
        mixed = cloud_api.ingest_latest(cloud_api.LatestBatch(site_id=identity["site_id"], readings=[
            record, dict(record, device_id=new_id),
        ]), identity)
        self.assertEqual(mixed["accepted"], 2)
        self.assertEqual(mixed["ignored_device_ids"], [self.device_id])
        self.assertEqual([row["device_id"] for row in self.client.get("/api/latest").json()["data"]], [new_id])

    def test_anonymous_viewer_and_other_customer_cannot_delete(self):
        with TestClient(cloud_api.app) as anonymous:
            self.assertEqual(anonymous.delete(self.url).status_code, 401)
        with cloud_api.connect() as db:
            db.execute("UPDATE customer_users SET role='viewer' WHERE id=%s", (self.user["id"],))
        self.assertEqual(self.client.delete(self.url).status_code, 403)
        other = seed_customer("Other tenant")
        self.client.post("/api/login", json={"username": other["user"]["customer_id"], "password": "local-delete-qa"})
        self.assertEqual(self.client.delete(self.url).status_code, 404)
        self.assertEqual(self.client.delete("/api/admin/devices/missing").status_code, 404)
        with cloud_api.connect() as db:
            self.assertTrue(db.execute("SELECT enabled FROM devices WHERE id=%s", (self.device_id,)).fetchone()[0])

    def test_audit_failure_rolls_back_device_and_revision(self):
        before = self.versions()
        with patch("cloud_api.record_configuration_audit", side_effect=RuntimeError("audit unavailable")):
            self.assertEqual(self.client.delete(self.url).status_code, 500)
        self.assertEqual(self.versions(), before)
        with cloud_api.connect() as db:
            self.assertEqual(db.execute("SELECT enabled,disabled_at FROM devices WHERE id=%s", (self.device_id,)).fetchone(), (True, None))

    def test_response_configuration_failure_rolls_back_deletion(self):
        before = self.versions()
        with patch("cloud_api.customer_configuration", side_effect=RuntimeError("read unavailable")):
            self.assertEqual(self.client.delete(self.url).status_code, 500)
        self.assertEqual(self.versions(), before)
        self.assertEqual(len(self.client.get("/api/config").json()["data"][0]["devices"]), 1)

    def test_config_snapshot_cannot_pair_old_devices_with_new_revision(self):
        identity = {"customer_id": self.user["customer_id"], "site_id": self.payload["site_id"]}
        old_revision = self.versions()[identity["site_id"]]
        original = cloud_api.customer_configuration
        paused_once = False

        def interleave(*args, **kwargs):
            nonlocal paused_once
            result = original(*args, **kwargs)
            if not paused_once:
                paused_once = True
                # Commit deletion after the reader has loaded the device list.
                with ThreadPoolExecutor(max_workers=1) as pool:
                    pool.submit(cloud_api.admin_delete_device, self.device_id, self.user).result(timeout=5)
            return result

        with patch("cloud_api.customer_configuration", side_effect=interleave):
            snapshot = cloud_api.edge_configuration_payload(identity)
        self.assertEqual(snapshot["revision"], old_revision)
        self.assertEqual(len(snapshot["rooms"][0]["devices"]), 1)
        next_snapshot = cloud_api.edge_configuration_payload(identity)
        self.assertEqual(next_snapshot["revision"], old_revision + 1)
        self.assertEqual(next_snapshot["rooms"][0]["devices"], [])

    def test_delete_waits_for_inflight_upload_and_concurrent_retries_are_idempotent(self):
        identity = {"customer_id": self.user["customer_id"], "site_id": self.payload["site_id"]}
        authorized, release = threading.Event(), threading.Event()
        original = cloud_api.authorize_ingest_target
        def pause(*args, **kwargs):
            names = original(*args, **kwargs)
            authorized.set()
            if not release.wait(5):
                raise RuntimeError("test synchronization timed out")
            return names
        reading = self.cached_reading(time.time())
        before = self.versions()
        with patch("cloud_api.authorize_ingest_target", side_effect=pause), ThreadPoolExecutor(max_workers=3) as pool:
            upload = pool.submit(cloud_api.ingest_latest, cloud_api.LatestBatch(site_id=identity["site_id"], readings=[reading]), identity)
            self.assertTrue(authorized.wait(3))
            deletions = [pool.submit(cloud_api.admin_delete_device, self.device_id, self.user) for _ in range(2)]
            try:
                # The explicitly held share lock is also observable from PostgreSQL.
                with cloud_api.connect() as db:
                    with self.assertRaises(Exception):
                        db.execute("SELECT 1 FROM customers WHERE id=%s FOR UPDATE NOWAIT", (identity["customer_id"],))
            finally:
                release.set()
            self.assertEqual(upload.result(timeout=5)["accepted"], 1)
            for deletion in deletions:
                self.assertTrue(deletion.result(timeout=5)["ok"])
        self.assertEqual(self.versions()[identity["site_id"]], before[identity["site_id"]] + 1)
        self.assertEqual(self.client.get("/api/latest").json()["data"], [])

    def cached_reading(self, measured_at):
        return dict(customer_id=self.user["customer_id"], site_id=self.payload["site_id"],
                    cleanroom_id=self.payload["cleanroom_id"], cleanroom_name="QA room",
                    device_id=self.device_id, device_name="QA device", measured_at=measured_at,
                    source="device", particles={}, environment={}, alarm_status="NORMAL")

    def test_offline_backlog_before_retirement_survives_but_new_readings_are_rejected(self):
        identity = {"customer_id": self.user["customer_id"], "site_id": self.payload["site_id"]}
        old_time = time.time() - 30
        with cloud_api.connect() as db:
            db.execute(
                """INSERT INTO email_alert_settings(customer_id,enabled,recipients,enabled_since)
                   VALUES(%s,true,'["qa@example.invalid"]'::jsonb,to_timestamp(%s))""",
                (identity["customer_id"], old_time - 60),
            )
        self.assertEqual(self.client.delete(self.url).status_code, 200)
        backlog = cloud_api.ReadingBatch(site_id=identity["site_id"], readings=[dict(self.cached_reading(old_time), record_uuid=uuid4())])
        self.assertEqual(len(cloud_api.ingest_batch(backlog, identity)["accepted"]), 1)
        event = dict(event_uuid=uuid4(), customer_id=identity["customer_id"], site_id=identity["site_id"],
                     cleanroom_id=self.payload["cleanroom_id"], device_id=self.device_id,
                     source="device", metric="temperature", started_at=old_time, limit_description="QA backlog")
        self.assertEqual(len(cloud_api.ingest_alarms(cloud_api.AlarmBatch(site_id=identity["site_id"], events=[event]), identity)["accepted"]), 1)
        with cloud_api.connect() as db:
            self.assertEqual(db.execute("SELECT count(*) FROM email_notifications WHERE customer_id=%s", (identity["customer_id"],)).fetchone()[0], 0)
        newer = cloud_api.ReadingBatch(site_id=identity["site_id"], readings=[dict(self.cached_reading(time.time()+1), record_uuid=uuid4())])
        with self.assertRaises(cloud_api.HTTPException) as rejected:
            cloud_api.ingest_batch(newer, identity)
        self.assertEqual(rejected.exception.status_code, 403)
        wrong_site = {**identity, "site_id": f'{identity["customer_id"]}-other'}
        with self.assertRaises(cloud_api.HTTPException):
            cloud_api.ingest_batch(cloud_api.ReadingBatch(site_id=wrong_site["site_id"], readings=[
                dict(self.cached_reading(old_time), site_id=wrong_site["site_id"], record_uuid=uuid4())
            ]), wrong_site)
        self.assertEqual(self.client.get("/api/latest").json()["data"], [])

    def test_collector_applies_removal_of_last_device(self):
        identity = {"customer_id": self.user["customer_id"], "site_id": self.payload["site_id"]}
        with tempfile.TemporaryDirectory() as directory:
            monitor = MonitoringService(Path(directory) / "collector.sqlite", lambda _: {}, lambda **_: {}, lambda *_: None)
            with patch.object(dashboard_server, "DB_PATH", monitor.db_path):
                dashboard_server.init_db()
            monitor.init_schema()
            monitor.apply_edge_configuration(cloud_api.edge_configuration_payload(identity))
            self.assertEqual(len(monitor.configuration(identity["customer_id"])[0]["devices"]), 1)
            self.assertEqual(self.client.delete(self.url).status_code, 200)
            updated = cloud_api.edge_configuration_payload(identity)
            self.assertEqual(updated["rooms"][0]["devices"], [])
            monitor.apply_edge_configuration(updated)
            self.assertEqual(monitor.configuration(identity["customer_id"])[0]["devices"], [])


if __name__ == "__main__":
    unittest.main()
