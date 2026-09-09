"""Lifecycle acceptance on the same explicitly opted-in disposable QA database."""
import io
import os
import time
import unittest
from unittest.mock import patch
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import load_workbook

import cloud_api
from email_alerts import device_email_suppressed, enqueue_alarm, claim_notification
import test_device_deletion_postgres as deletion_tests
from test_device_deletion_postgres import seed_customer


@unittest.skipUnless(os.environ.get('DCP_DEVICE_DELETE_TEST_DATABASE_URL'), 'Opt-in isolated PostgreSQL required')
class DeviceLifecyclePostgresTests(unittest.TestCase):
    setUp = deletion_tests.DeviceDeletionPostgresTests.setUp
    tearDown = deletion_tests.DeviceDeletionPostgresTests.tearDown
    versions = deletion_tests.DeviceDeletionPostgresTests.versions

    @classmethod
    def setUpClass(cls):
        deletion_tests.DeviceDeletionPostgresTests.setUpClass.__func__(cls)

    @classmethod
    def tearDownClass(cls):
        deletion_tests.DeviceDeletionPostgresTests.tearDownClass.__func__(cls)

    def item(self):
        response=self.client.get('/api/admin/device-lifecycle')
        self.assertEqual(response.status_code,200,response.text)
        return next(d for d in response.json()['data']['devices'] if d['id']==self.device_id)

    def edit(self, **changes):
        current=self.item()
        return self.client.patch(self.url,json={'name':current['name'],'cleanroom_id':current['cleanroom_id'],
                                               'expected_updated_at':current['updated_at'],**changes})

    def reading(self, stamp=None, room=None):
        return dict(customer_id=self.user['customer_id'],site_id=self.payload['site_id'],
                    cleanroom_id=room or self.payload['cleanroom_id'],cleanroom_name='untrusted',device_id=self.device_id,
                    device_name='untrusted',measured_at=stamp or time.time(),source='device',particles={},environment={},alarm_status='normal')

    def test_manual_registration_preserves_nondefault_port_in_edge_configuration(self):
        payload = dict(self.payload, name='Gateway port 15020', tcp_port=15020)
        response = self.client.post('/api/admin/devices', json=payload)
        self.assertEqual(response.status_code, 201, response.text)
        rooms = cloud_api.edge_configuration_payload({
            'customer_id': self.user['customer_id'], 'site_id': self.payload['site_id'],
        })['rooms']
        device = next(d for room in rooms for d in room['devices'] if d['name'] == payload['name'])
        self.assertEqual(device['tcpPort'], 15020)
        # Same IP and slave on a different port is a separate endpoint;
        # repeating the exact endpoint must still be rejected.
        self.assertNotEqual(device['id'], self.device_id)
        self.assertEqual(self.client.post('/api/admin/devices', json=payload).status_code, 409)

    def ingest(self, record):
        identity={'customer_id':self.user['customer_id'],'site_id':record['site_id']}
        return cloud_api.ingest_batch(cloud_api.ReadingBatch(site_id=record['site_id'],readings=[dict(record,record_uuid=uuid4())]),identity)

    def test_edit_preserves_history_and_old_backlog_but_rejects_stale_writes(self):
        old=self.reading();self.ingest(old)
        revision=self.item()['updated_at']
        new_room=cloud_api.create_cloud_cleanroom(self.user['customer_id'],'Moved room',actor_kind='customer_user',actor_user_id=self.user['id'])
        response=self.edit(name='Renamed instrument',cleanroom_id=new_room,host='192.168.50.31')
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.item()['id'],self.device_id)
        # Old workshop remains selectable and its original labels survive.
        rooms=self.client.get('/api/config?include_disabled=true').json()['data']
        self.assertTrue(any(d['id']==self.device_id for r in rooms if r['id']==self.payload['cleanroom_id'] for d in r['devices']))
        self.ingest(old)
        rows=self.client.get('/api/history',params={'cleanroom_ids':self.payload['cleanroom_id'],'device_ids':self.device_id}).json()['data']
        self.assertEqual(len(rows),2);self.assertTrue(all(r['device']==self.payload['name'] for r in rows))
        workbook=load_workbook(io.BytesIO(self.client.get('/api/history/export-xlsx',params={'cleanroom_ids':self.payload['cleanroom_id'],'device_ids':self.device_id}).content))
        self.assertEqual(len(workbook.sheetnames),2)
        self.assertEqual(workbook.worksheets[1].cell(2,2).value,self.payload['name'])
        with self.assertRaises(HTTPException):self.ingest(self.reading())
        self.assertEqual(self.client.patch(self.url,json={'name':'Stale','cleanroom_id':new_room,'expected_updated_at':revision}).status_code,409)
        self.assertEqual(self.edit(name='Renamed instrument').status_code,200)
        logs=self.client.get('/api/admin/device-audit',params={'device_id':self.device_id}).json()['data']['items']
        edit=next(a for a in logs if a['action']=='device.updated')
        self.assertEqual(edit['details']['before']['name'],self.payload['name'])
        self.assertEqual(edit['details']['after']['cleanroom_id'],new_room)

    def test_restore_conflict_idempotence_capacity_and_gap(self):
        self.assertEqual(self.client.delete(self.url).status_code,200)
        gap=self.reading()
        other=cloud_api.create_cloud_device(self.user['customer_id'],dict(self.payload,name='Conflicting instrument'),actor_kind='customer_user',actor_user_id=self.user['id'])
        self.assertEqual(self.client.post(self.url+'/restore').status_code,409)
        self.assertFalse(self.item()['enabled'])
        self.client.delete('/api/admin/devices/'+other)
        with patch.object(cloud_api,'MAX_ACTIVE_DEVICES',0):self.assertEqual(self.client.post(self.url+'/restore').status_code,409)
        self.assertEqual(self.client.post(self.url+'/restore').status_code,200)
        versions=self.versions();self.assertEqual(self.client.post(self.url+'/restore').status_code,200);self.assertEqual(versions,self.versions())
        with self.assertRaises(HTTPException):self.ingest(gap)
        self.ingest(self.reading())
        identity={'customer_id':self.user['customer_id'],'site_id':self.payload['site_id']}
        result=cloud_api.ingest_latest(cloud_api.LatestBatch(site_id=identity['site_id'],readings=[gap]),identity)
        self.assertEqual(result['ignored_device_ids'],[self.device_id])

    def test_permissions_and_cross_tenant_are_enforced_for_all_operations(self):
        current=self.item(); other=seed_customer('Other tenant')
        other_url='/api/admin/devices/'+other['device_id']
        body={'name':'Bad','cleanroom_id':current['cleanroom_id'],'expected_updated_at':current['updated_at']}
        self.assertEqual(self.client.patch(other_url,json=body).status_code,404)
        self.assertEqual(self.client.post(other_url+'/restore').status_code,404)
        self.assertEqual(self.client.post(other_url+'/maintenance',json={'reason':'QA','ends_at':time.time()+3600}).status_code,404)
        self.assertEqual(self.client.delete(other_url+'/maintenance').status_code,404)
        self.assertEqual(self.client.get('/api/admin/device-audit',params={'device_id':other['device_id']}).json()['data']['items'],[])
        self.assertEqual(self.edit(cleanroom_id=other['payload']['cleanroom_id']).status_code,404)
        with cloud_api.connect() as db:db.execute("UPDATE customer_users SET role='viewer' WHERE id=%s",(self.user['id'],))
        for method,url,payload in [('get','/api/admin/device-lifecycle',None),('get','/api/admin/device-audit',None),('patch',self.url,body),('post',self.url+'/restore',None),('post',self.url+'/maintenance',{'reason':'QA','ends_at':time.time()+3600}),('delete',self.url+'/maintenance',None)]:
            self.assertEqual(self.client.request(method,url,**({'json':payload} if payload else {})).status_code,403)

    def test_legacy_manager_cannot_change_connections(self):
        with cloud_api.connect() as db:db.execute("UPDATE customer_users SET role='customer' WHERE id=%s",(self.user['id'],))
        self.assertEqual(self.edit(name='Allowed rename').status_code,200)
        self.assertEqual(self.edit(host='192.168.50.99').status_code,403)

    def test_failed_audit_rolls_back_edit_restore_and_maintenance(self):
        initial=self.item();versions=self.versions()
        with patch.object(cloud_api,'record_configuration_audit',side_effect=RuntimeError('QA rollback')):
            self.assertEqual(self.edit(name='Rollback').status_code,500)
            self.assertEqual(self.client.post(self.url+'/maintenance',json={'reason':'Rollback','ends_at':time.time()+3600}).status_code,500)
        self.assertEqual(self.item()['name'],initial['name']);self.assertEqual(self.versions(),versions)
        self.assertIsNone(self.item()['maintenance_until'])
        self.client.delete(self.url)
        with patch.object(cloud_api,'record_configuration_audit',side_effect=RuntimeError('QA rollback')):self.assertEqual(self.client.post(self.url+'/restore').status_code,500)
        self.assertFalse(self.item()['enabled'])

    def alarm(self, stamp):
        event=uuid4();identity={'customer_id':self.user['customer_id'],'site_id':self.payload['site_id']}
        cloud_api.ingest_alarms(cloud_api.AlarmBatch(site_id=identity['site_id'],events=[dict(event_uuid=event,**identity,cleanroom_id=self.payload['cleanroom_id'],device_id=self.device_id,source='device',metric='temperature',started_at=stamp,limit_description='QA')]),identity)
        return event

    def test_maintenance_preserves_data_suppresses_email_and_ends_without_backfill(self):
        # Enable only the local persisted policy, never an SMTP transport.
        with cloud_api.connect() as db:
            db.execute("INSERT INTO email_alert_settings(customer_id,enabled,recipients,notify_recovery,enabled_since) VALUES(%s,true,'[\"qa@example.invalid\"]',true,now()-interval '1 day')",(self.user['customer_id'],))
        before_event=self.alarm(time.time())
        response=self.client.post(self.url+'/maintenance',json={'reason':'Calibration QA','ends_at':time.time()+3600})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.item()['maintenance_reason'],'Calibration QA')
        during=time.time();self.ingest(self.reading(during)); event=self.alarm(during)
        with cloud_api.connect() as db:
            self.assertTrue(device_email_suppressed(db,self.user['customer_id'],event))
            self.assertEqual(db.execute('SELECT count(*) FROM email_notifications WHERE event_uuid=%s',(event,)).fetchone()[0],0)
            self.assertEqual(db.execute('SELECT status FROM email_notifications WHERE event_uuid=%s',(before_event,)).fetchone()[0],'cancelled')
        self.assertEqual(self.client.get('/api/history').json()['data'][0]['maintenance']['reason'],'Calibration QA')
        workbook=load_workbook(io.BytesIO(self.client.get('/api/history/export-xlsx').content))
        self.assertIn('Calibration QA',workbook.worksheets[1].cell(2,12).value)
        self.assertEqual(self.client.delete(self.url+'/maintenance').status_code,200)
        self.assertIsNone(self.item()['maintenance_until'])
        with cloud_api.connect() as db:self.assertTrue(device_email_suppressed(db,self.user['customer_id'],event))
        fresh=self.alarm(time.time())
        with cloud_api.connect() as db:self.assertFalse(device_email_suppressed(db,self.user['customer_id'],fresh))

    def test_maintenance_validation_expiry_and_retirement(self):
        for reason,end in [('',time.time()+3600),('bad',time.time()-1),('bad',time.time()+8*86400)]:
            self.assertIn(self.client.post(self.url+'/maintenance',json={'reason':reason,'ends_at':end}).status_code,(400,422))
        self.assertEqual(self.client.post(self.url+'/maintenance',json={'reason':'Expiry QA','ends_at':time.time()+3600}).status_code,200)
        self.assertEqual(self.client.post(self.url+'/maintenance',json={'reason':'Duplicate','ends_at':time.time()+3600}).status_code,409)
        with cloud_api.connect() as db:db.execute("UPDATE device_maintenance_periods SET started_at=now()-interval '2 hours',ends_at=now()-interval '1 hour' WHERE device_id=%s",(self.device_id,))
        self.assertIsNone(self.item()['maintenance_until'])
        self.assertEqual(self.client.post(self.url+'/maintenance',json={'reason':'Second','ends_at':time.time()+3600}).status_code,200)
        self.client.delete(self.url);self.client.post(self.url+'/restore')
        self.assertIsNone(self.item()['maintenance_until'])

    def test_site_move_advances_both_collectors_and_sync_does_not_claim_readings(self):
        before=self.versions();new_site=self.user['customer_id']+'-other'
        self.assertEqual(self.edit(site_id=new_site).status_code,200)
        after=self.versions()
        self.assertEqual(after[new_site],before[new_site]+1)
        self.assertEqual(after[self.payload['site_id']],before[self.payload['site_id']]+1)
        self.assertEqual(self.item()['sync_state'],'pending')
        with cloud_api.connect() as db:
            db.execute("UPDATE sites SET applied_config_version=config_version,config_apply_status='applied',last_heartbeat_at=now() WHERE id=%s",(new_site,))
        self.assertEqual(self.item()['sync_state'],'applied')
        self.assertIsNone(self.item()['last_seen_at'])
        with cloud_api.connect() as db:db.execute("UPDATE sites SET config_apply_status='failed',config_apply_error='QA failure' WHERE id=%s",(new_site,))
        self.assertEqual(self.item()['sync_state'],'failed')

    def test_concurrent_edit_conflicts_and_schema_reapply_preserves_assignment_periods(self):
        from concurrent.futures import ThreadPoolExecutor
        data=self.item();payload={'name':'Concurrent','cleanroom_id':data['cleanroom_id'],'expected_updated_at':data['updated_at']}
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _:self.client.patch(self.url,json=payload).status_code,range(2)))
        self.assertEqual(sorted(results),[200,409])
        with cloud_api.connect() as db:
            before=db.execute('SELECT count(*) FROM device_assignment_periods WHERE device_id=%s',(self.device_id,)).fetchone()[0]
        cloud_api.initialize_schema()
        with cloud_api.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM device_assignment_periods WHERE device_id=%s',(self.device_id,)).fetchone()[0],before)

    def test_legacy_settings_rename_is_also_audited(self):
        config=self.client.get('/api/config').json()['data']
        config[0]['devices'][0]['name']='Renamed in settings'
        result=self.client.post('/api/config',json={'rooms':config})
        self.assertEqual(result.status_code,200,result.text)
        items=self.client.get('/api/admin/device-audit',params={'device_id':self.device_id}).json()['data']['items']
        self.assertEqual(items[0]['action'],'device.updated')
        self.assertEqual(items[0]['details']['before']['name'],self.payload['name'])
        self.assertEqual(items[0]['details']['after']['name'],'Renamed in settings')
